#!/usr/bin/env python3
"""Make a lightweight raster preview when no office renderer is available.

This is a QA fallback, not a PowerPoint-compatible rendering authority.  It
places the exact embedded figures and uses Noto Sans CJK metrics for text-box
overflow inspection.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[4]
PPTX = ROOT / "PPT0821/ppt_SH3_L2_five_family_solar_20260821.pptx"
OUTPUT = Path("/tmp/ppt_SH3_L2_five_family_solar_preview")
REGULAR = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
BOLD = "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"
CANVAS = (1600, 900)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, width: int) -> list[str]:
    if not text:
        return [""]
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if current and draw.textlength(candidate, font=font) > width:
            lines.append(current)
            current = char
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def main() -> int:
    deck = Presentation(PPTX)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    sx = CANVAS[0] / deck.slide_width
    sy = CANVAS[1] / deck.slide_height
    for slide_index, slide in enumerate(deck.slides, 1):
        canvas = Image.new("RGB", CANVAS, "white")
        draw = ImageDraw.Draw(canvas)
        for shape in slide.shapes:
            left = int(round(shape.left * sx))
            top = int(round(shape.top * sy))
            width = max(1, int(round(shape.width * sx)))
            height = max(1, int(round(shape.height * sy)))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = Image.open(BytesIO(shape.image.blob)).convert("RGB")
                image = image.resize((width, height), Image.Resampling.LANCZOS)
                canvas.paste(image, (left, top))
                continue
            if not getattr(shape, "has_text_frame", False) or not shape.text:
                continue
            margin_left = int(round(shape.text_frame.margin_left * sx))
            margin_right = int(round(shape.text_frame.margin_right * sx))
            margin_top = int(round(shape.text_frame.margin_top * sy))
            x = left + margin_left
            y = top + margin_top
            usable_width = max(1, width - margin_left - margin_right)
            for paragraph in shape.text_frame.paragraphs:
                run = paragraph.runs[0] if paragraph.runs else None
                size_pt = float(run.font.size.pt) if run is not None and run.font.size else 18.0
                size_px = max(8, int(round(size_pt * CANVAS[0] / 960.0)))
                bold = bool(run is not None and run.font.bold)
                font = ImageFont.truetype(BOLD if bold else REGULAR, size_px)
                for line in wrap_text(draw, paragraph.text, font, usable_width):
                    draw.text((x, y), line, font=font, fill="#222222")
                    y += int(round(size_px * 1.28))
                y += max(1, int(round(size_px * 0.10)))
            if y > top + height:
                raise RuntimeError(
                    f"Approximate text overflow on slide {slide_index}, shape {shape.name}: "
                    f"{y - (top + height)} px"
                )
        canvas.save(OUTPUT / f"slide_{slide_index:02d}.png")
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
