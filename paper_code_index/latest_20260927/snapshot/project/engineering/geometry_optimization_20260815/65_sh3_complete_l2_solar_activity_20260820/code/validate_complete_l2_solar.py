#!/usr/bin/env python3
"""Validate the five-family quiet-L2 source set, SH3 projection, and PPT."""

from __future__ import annotations

import csv
import json
import math
import zipfile
from collections import defaultdict
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[4]
PACKAGE = Path(__file__).resolve().parents[1]
TABLES = PACKAGE / "outputs/tables"
FIGURES = PACKAGE / "outputs/figures"
SUMMARY = PACKAGE / "outputs/summary.json"
PPTX = ROOT / "PPT0821/ppt_SH3_L2_five_family_solar_20260821.pptx"
VALIDATION = PACKAGE / "data/validation.json"

CENTRAL = "sun_earth_l2_quiet_1au_proxy"
SOLAR_MAX = "sun_earth_l2_solar_max_2014"
FAMILIES = ("gamma", "p", "alpha", "eminus", "eplus")
EXPECTED_SUPPORT_KEV = {
    "gamma": (3.0, 739072203.3525774),
    "p": (5000.0, 1.0e10),
    "alpha": (10000.0, 1.0e10),
    "eminus": (1000.0, 1.0e10),
    "eplus": (1000.0, 1.0e10),
}


def rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def grouped_curves(path: Path) -> dict[str, list[tuple[float, float]]]:
    grouped: dict[str, list[tuple[float, float]]] = defaultdict(list)
    for row in rows(path):
        assert row["angular_domain"] == "full_sphere_4pi"
        assert math.isclose(float(row["solid_angle_sr"]), 4.0 * math.pi, rel_tol=1.0e-12)
        energy = float(row["energy_keV_total"])
        flux = float(row["differential_flux_cm2_s_keV"])
        assert energy > 0 and flux > 0 and math.isfinite(energy) and math.isfinite(flux)
        grouped[row["family"]].append((energy, flux))
    assert set(grouped) == set(FAMILIES)
    for family, curve in grouped.items():
        energies = [item[0] for item in curve]
        assert all(b > a for a, b in zip(energies, energies[1:])), family
        expected_min, expected_max = EXPECTED_SUPPORT_KEV[family]
        assert math.isclose(energies[0], expected_min, rel_tol=2.0e-12)
        assert math.isclose(energies[-1], expected_max, rel_tol=2.0e-12)
    return grouped


def nearest(curve: list[tuple[float, float]], energy_keV: float) -> float:
    return min(curve, key=lambda item: abs(math.log(item[0] / energy_keV)))[1]


def main() -> int:
    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    assert summary["status"] == "PASS__SH3_FIVE_FAMILY_CONTINUOUS_QUIET_L2_SOURCE_SET_AND_SOLAR_ACTIVITY_PROXY"
    assert "NOT_COMPLETE_L2_ENVIRONMENT" in summary["authority_boundary"]

    central = grouped_curves(TABLES / "l2_source_spectra.csv")
    solar_max = grouped_curves(TABLES / "l2_solar_max_source_spectra.csv")
    for family in FAMILIES:
        assert len(central[family]) == len(solar_max[family])

    for energy_gev in (1.366, 2.0, 10.0, 48.4):
        energy_keV = energy_gev * 1.0e6
        assert nearest(solar_max["p"], energy_keV) < nearest(central["p"], energy_keV)
    assert nearest(solar_max["alpha"], 20.0e6) < nearest(central["alpha"], 20.0e6)

    by_env = summary["environment_estimates"]
    f_full = float(by_env[CENTRAL]["estimated_20d_F3_ph_cm2_s"])
    f_solar_max = float(by_env[SOLAR_MAX]["estimated_20d_F3_ph_cm2_s"])
    assert f_full > f_solar_max > 0
    assert math.isclose(f_full, 4.8782424044738405e-05, rel_tol=2.0e-12)
    assert math.isclose(f_solar_max, 3.4459876347906416e-05, rel_tol=2.0e-12)

    dominant = summary["dominant_L2_activation_key"]
    assert dominant["shared_delayed_survivors"] == 4
    assert dominant["activation_RP_records"] == 1
    assert math.isclose(float(dominant["primary_energy_MeV"]), 1365.984184, rel_tol=1.0e-12)
    assert float(dominant["fraction_of_total_L2_background"]) > 0.88
    assert float(summary["eplus_coverage_proxy_impact"]["fraction_of_total_L2_background"]) < 1.0e-5

    response = {row["family"]: row for row in rows(TABLES / "sh3_response_weighted_primary_energy_bands.csv")}
    assert 320.0 < float(response["alpha"]["response_weighted_primary_energy_p10_MeV"])
    assert float(response["alpha"]["response_weighted_primary_energy_p90_MeV"]) < 400000.0
    stream_response = {
        (row["stream"], row["family"]): row
        for row in rows(TABLES / "sh3_prompt_delayed_response_weighted_primary_energy_bands.csv")
    }
    assert set(stream_response) == {
        ("prompt", "gamma"),
        ("delayed", "gamma"),
        ("delayed", "eplus"),
        ("delayed", "p"),
        ("delayed", "alpha"),
        ("delayed", "n"),
    }
    assert int(stream_response[("prompt", "gamma")]["selected_W2_final_events"]) == 12
    assert int(stream_response[("delayed", "gamma")]["selected_W2_final_events"]) == 52
    assert math.isclose(
        float(stream_response[("prompt", "gamma")]["response_weighted_primary_energy_p10_MeV"]),
        0.598334,
        rel_tol=2.0e-6,
    )
    assert float(stream_response[("delayed", "gamma")]["response_weighted_primary_energy_p90_MeV"]) > 6.0e5

    soft = rows(TABLES / "l2_directional_soft_proton_spectra.csv")
    assert len(soft) == 360
    assert all(row["performance_role"].startswith("SEPARATE_DIRECTIONAL_SOURCE") for row in soft)

    for name in (
        "01_background_energy_bands_sh3_l2.png",
        "02_full_spectrum_comparison_sh3_l2.png",
        "03_normalized_minimum_detectable_flux_sh3_l2.png",
        "04_l2_solar_activity_sh3.png",
    ):
        assert (FIGURES / name).is_file()

    with zipfile.ZipFile(PPTX) as archive:
        assert archive.testzip() is None
    deck = Presentation(PPTX)
    assert len(deck.slides) == 6
    slide_text = ["\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")) for slide in deck.slides]
    assert "20260821" in slide_text[0]
    assert "HZE" in slide_text[1]
    assert "1.366 GeV" in slide_text[3]
    assert "prompt 蓝；delayed 橙" in slide_text[2]
    assert "2009 full" in slide_text[4] and "2014" in slide_text[4]
    assert "Athena 80%" not in "\n".join(slide_text)

    forbidden = "cosima_spectra_dp_2602" + "units"
    text_files = [
        path
        for path in PACKAGE.rglob("*")
        if path.is_file() and path.suffix.lower() in {".py", ".md", ".csv", ".json", ".txt"}
    ]
    assert all(forbidden not in path.read_text(encoding="utf-8", errors="ignore") for path in text_files)

    result = {
        "status": "PASS__PHYSICS_CONTRACT_TABLES_PPT_STRUCTURE_AND_LIGHTWEIGHT_PREVIEW",
        "office_renderer": "UNAVAILABLE__LIBREOFFICE_INSTALL_AND_USERSPACE_DOWNLOAD_FAILED",
        "preview_qa": {
            "method": "exact embedded figures plus Noto Sans CJK text metrics",
            "location": "/tmp/ppt_SH3_L2_five_family_solar_preview",
            "text_overflow": 0,
            "visually_inspected_slides": [2, 3, 4, 5],
        },
        "spectral_families": list(FAMILIES),
        "scenario_Fmin_ph_cm2_s": {
            "2009_full_GCR_max": f_full,
            "2014_GCR_min": f_solar_max,
        },
        "dominant_activation_key": dominant,
        "limitations_asserted": [
            "five-family quiet source set is not a complete L2 physical environment",
            "HZE is a missing static GCR component",
            "SEP, directional soft protons and Galactic diffuse gamma remain separate",
            "conditional MC error excludes shared single-RP correlation and source-model uncertainty",
        ],
    }
    VALIDATION.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
