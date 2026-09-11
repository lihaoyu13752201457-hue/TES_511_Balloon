#!/usr/bin/env python3
"""Build the editable seven-slide TES-511 source/workflow briefing."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "TES511_balloon_to_LEO_source_compute_workflow_brief_20260813.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Noto Sans CJK SC"

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(23, 33, 43)
MUTED = RGBColor(83, 96, 109)
DIM = RGBColor(116, 128, 139)
BLUE = RGBColor(7, 91, 157)
BLUE_SOFT = RGBColor(234, 244, 251)
GREEN = RGBColor(35, 122, 87)
GREEN_SOFT = RGBColor(234, 246, 240)
AMBER = RGBColor(170, 103, 8)
AMBER_SOFT = RGBColor(255, 245, 223)
PANEL = RGBColor(246, 248, 250)
BORDER = RGBColor(224, 230, 235)
HEADER_BLUE = RGBColor(9, 98, 163)


NOTES = [
    (
        "从气球到近地轨道",
        "这份简报只回答三个问题：气球和近地轨道的入射环境有什么不同，完整任务会带来多少额外计算维度，以及我们现在怎样组织计算、存储和数据处理。",
    ),
    (
        "入射组件谱",
        "先看入射源本身。气球环境受残余大气层影响，轨道环境则要区分天区、地球反照、原初和次级粒子，所以不能把一套谱简单乘一个系数后当作另一套环境。图中不包含活化和延迟分量，也不是探测器计数谱。",
    ),
    (
        "关键能带",
        "正常轨道的宽带伽马在部分能区更低，但质子和正电子在一些能段更强。450到600 keV的比值只是宽能带代理，不能替代窄511 keV线及探测器响应计算。",
    ),
    (
        "计算负担",
        "如果只比较一个固定状态，气球和轨道的prompt输运可以按同一数量级规划。真正抬高轨道任务成本的是状态维度和完整物理链，因此先用代表状态试跑，测量每个primary的墙钟、内存和输出量。",
    ),
    (
        "当前计算方式",
        "当前流程先冻结物理契约，再把大任务拆成短而可恢复的分片。分片规模根据实测耗时自适应调整，重尾粒子单独调度，每一片完成后立即验收，所以失败只需要补缺失的小片。",
    ),
    (
        "存储与数据处理",
        "原始层、验证层和分析层分开保存。原始产物继续留在runs目录，receipt记录可审计证据，日常分析使用轻量派生输出，从而减少重复读取和重复解析。77.6 GB和62.8 MB不是原始文件的前后压缩关系。",
    ),
    (
        "早期与当前流程",
        "当前流程的核心变化是短分片、即时验证、分层存储和严格分域。它缩小失败后的重算范围，也让后续响应和任务参数变化不必反复扫描全部原始文件。轨道版本下一步应先做代表状态pilot。",
    ),
]


def set_font(run, size: float, color: RGBColor, *, bold: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run._r.get_or_add_rPr().set(qn("a:ea"), FONT)


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float = 16,
    color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
    space_after: float = 0.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = line
        set_font(run, size, color, bold=bold)
    return shape


def add_title(slide, section: str, title: str, *, title_size: float = 28) -> None:
    add_text(slide, 0.62, 0.44, 12.0, 0.22, section, size=9, color=BLUE, bold=True)
    add_text(slide, 0.62, 0.72, 12.0, 0.56, title, size=title_size, color=INK, bold=True)


def add_chrome(slide, index: int, total: int) -> None:
    base = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(0.035)
    )
    base.fill.solid()
    base.fill.fore_color.rgb = RGBColor(232, 237, 242)
    base.line.fill.background()
    progress = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        Inches(SLIDE_W * (index + 1) / total),
        Inches(0.035),
    )
    progress.fill.solid()
    progress.fill.fore_color.rgb = BLUE
    progress.line.fill.background()
    add_text(
        slide,
        12.16,
        7.13,
        0.57,
        0.17,
        f"{index + 1} / {total}",
        size=7.5,
        color=DIM,
        align=PP_ALIGN.RIGHT,
    )


def add_round_rect(slide, x, y, w, h, fill: RGBColor, *, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.7)
    return shape


def add_picture_fit(slide, path: Path, x: float, y: float, w: float, h: float) -> None:
    with Image.open(path) as image:
        iw, ih = image.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px = x + (w - pw) / 2
    py = y + (h - ph) / 2
    slide.shapes.add_picture(
        str(path), Inches(px), Inches(py), width=Inches(pw), height=Inches(ph)
    )


def add_pill(slide, x: float, y: float, w: float, text: str, *, fill=PANEL, color=MUTED):
    add_round_rect(slide, x, y, w, 0.34, fill)
    add_text(
        slide,
        x + 0.10,
        y + 0.075,
        w - 0.20,
        0.18,
        text,
        size=8.5,
        color=color,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_notes(slide, title: str, script: str) -> None:
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"{title}\n\n{script}"


def build_cover(slide) -> None:
    add_text(
        slide,
        1.0,
        1.70,
        11.33,
        0.25,
        "MASS_MODEL_511 · S3D-O8",
        size=9,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        1.0,
        2.05,
        11.33,
        1.25,
        "从 38 km 气球\n到 530 km 近地轨道",
        size=36,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    rule = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(5.84),
        Inches(3.55),
        Inches(1.65),
        Inches(0.06),
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = BLUE
    rule.line.fill.background()
    add_text(
        slide,
        1.0,
        3.88,
        11.33,
        0.35,
        "入射源环境、计算负担与当前模拟流程",
        size=16,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_pill(slide, 5.28, 4.58, 2.78, "简化技术报告 · 2026-08-13", fill=BLUE_SOFT, color=BLUE)


def build_source_slide(slide, *, title: str, image: str, caption: str, section: str) -> None:
    add_title(slide, section, title, title_size=27)
    panel = add_round_rect(slide, 0.72, 1.38, 11.89, 4.78, WHITE, line=BORDER)
    panel.shadow.inherit = False
    add_picture_fit(slide, ASSETS / image, 0.88, 1.55, 11.57, 4.44)
    add_text(
        slide,
        1.00,
        6.28,
        11.33,
        0.34,
        caption,
        size=10.2,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_pill(slide, 4.43, 6.70, 4.47, "入射源级 · 不含活化 · 不是探测器计数", fill=PANEL, color=MUTED)


def add_burden_card(slide, x, fill, kicker, main, detail) -> None:
    add_round_rect(slide, x, 2.38, 3.72, 3.10, fill)
    add_text(slide, x + 0.23, 2.67, 3.25, 0.23, kicker, size=9, color=MUTED, bold=True)
    add_text(slide, x + 0.23, 3.06, 3.25, 0.72, main, size=21, color=INK, bold=True)
    add_text(slide, x + 0.23, 4.18, 3.25, 0.84, detail, size=11, color=MUTED)


def build_burden(slide) -> None:
    add_title(slide, "02 · 计算负担", "单状态近似同量级，完整轨道链更重")
    add_text(
        slide,
        1.0,
        1.50,
        11.33,
        0.34,
        "计算规模 ≈ 状态数 × 方向数 × 粒子族 × 物理链",
        size=17,
        color=BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_burden_card(
        slide,
        0.62,
        BLUE_SOFT,
        "代表状态",
        "静态 prompt\n可按同量级预算",
        "单个轨道状态与单个气球状态，都可用相同的分片输运框架完成。",
    )
    add_burden_card(
        slide,
        4.81,
        AMBER_SOFT,
        "完整任务",
        "轨道状态维度\n显著增加",
        "轨道位置、指向、暴露权重，以及高辐射区后的活化与延迟链均需审计。",
    )
    add_burden_card(
        slide,
        9.00,
        PANEL,
        "推荐做法",
        "先做代表状态\npilot 再扩展",
        "用短试跑测量每族耗时、输出体积和方差，再决定轨道分箱与统计量。",
    )
    add_text(
        slide,
        1.0,
        5.85,
        11.33,
        0.30,
        "当前分片主线：6,149,560 validated primaries · 488 validated receipts",
        size=11,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        1.0,
        6.30,
        11.33,
        0.38,
        "当前没有可靠的轨道全链成本倍率；先 pilot，再按实测资源扩展。",
        size=12.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_pipeline_box(slide, x, fill, number, title, body) -> None:
    add_round_rect(slide, x, 2.05, 2.62, 3.48, fill)
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x + 0.20), Inches(2.28), Inches(0.43), Inches(0.43)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLUE
    circle.line.fill.background()
    add_text(
        slide,
        x + 0.20,
        2.36,
        0.43,
        0.16,
        str(number),
        size=8.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(slide, x + 0.20, 2.91, 2.18, 0.54, title, size=17.5, color=INK, bold=True)
    add_text(slide, x + 0.20, 3.68, 2.18, 1.20, body, size=10.5, color=MUTED)


def build_pipeline(slide) -> None:
    add_title(slide, "03 · 当前计算方式", "物理条件冻结，计算切成可恢复短分片")
    xs = [0.60, 3.65, 6.70, 9.75]
    fills = [BLUE_SOFT, PANEL, GREEN_SOFT, PANEL]
    contents = [
        ("冻结计算契约", "源、几何、模式、粒子族和随机种子均显式登记。"),
        ("短分片自适应", "目标约 5–10 分钟/任务；依据近期实测耗时调整下一片规模。"),
        ("按粒子族调度", "高成本、重尾粒子使用更小批次与轮转排程。"),
        ("分片即刻验收", "每片生成 receipt；只在同几何、模式和粒子族域内汇总。"),
    ]
    for idx, (x, fill, content) in enumerate(zip(xs, fills, contents), start=1):
        add_pipeline_box(slide, x, fill, idx, content[0], content[1])
        if idx < 4:
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(x + 2.69),
                Inches(3.78),
                Inches(x + 2.98),
                Inches(3.78),
            )
            connector.line.color.rgb = RGBColor(141, 161, 177)
            connector.line.width = Pt(1.6)
            connector.line.end_arrowhead = True
    add_text(
        slide,
        0.90,
        5.92,
        11.53,
        0.34,
        "Batch0006：6.149M validated primaries · 488 validated receipts",
        size=12,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        0.90,
        6.38,
        11.53,
        0.34,
        "优化任务组织、并发和 I/O；同一批次内物理输入保持不变。",
        size=11.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_metric(slide, x, value, label) -> None:
    add_round_rect(slide, x, 1.55, 2.82, 1.02, PANEL)
    add_text(slide, x + 0.12, 1.72, 2.58, 0.33, value, size=21, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, x + 0.12, 2.13, 2.58, 0.24, label, size=8.7, color=MUTED, align=PP_ALIGN.CENTER)


def add_storage_layer(slide, x, fill, title, body) -> None:
    add_round_rect(slide, x, 3.05, 3.80, 2.32, fill)
    add_text(slide, x + 0.22, 3.35, 3.36, 0.37, title, size=15, color=INK, bold=True)
    add_text(slide, x + 0.22, 3.98, 3.36, 0.92, body, size=10.5, color=MUTED)


def build_storage(slide) -> None:
    add_title(slide, "03 · 存储与数据处理", "原始层保留，验证与分析走轻量派生层")
    metrics = [
        (0.64, "976", "selected jobs"),
        (3.64, "13,777,092", "primary histories"),
        (6.64, "77.6 GB", "976 selected .sim.gz"),
        (9.64, "62.8 MB", "当前活跃派生输出"),
    ]
    for metric in metrics:
        add_metric(slide, *metric)
    add_storage_layer(
        slide,
        0.64,
        BLUE_SOFT,
        "① 原始输运层",
        "SIM / DAT / log 保留完整信息；临时文件完成后原子发布，失败尝试单独隔离。",
    )
    add_storage_layer(
        slide,
        4.77,
        PANEL,
        "② receipt / ledger",
        "每片记录种子、histories、TT/RP、耗时、内存、字节数与校验摘要。",
    )
    add_storage_layer(
        slide,
        8.90,
        GREEN_SOFT,
        "③ 派生分析层",
        "一次解析形成 compact catalogs；响应、筛选和任务折叠复用轻量表格。",
    )
    add_text(
        slide,
        0.80,
        5.73,
        11.73,
        0.38,
        "34 个 compact catalogs：16 prompt + 16 delayed + 2 signal",
        size=11.5,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        0.80,
        6.20,
        11.73,
        0.47,
        "77.6 GB 与 62.8 MB 表示输运层和派生工作层分开保存；原始数据没有被替代或删除。",
        size=10.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def set_cell(cell, text: str, *, size: float, color: RGBColor, bold=False, align=PP_ALIGN.LEFT, fill=WHITE) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.10)
    cell.margin_right = Inches(0.10)
    cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, color, bold=bold)


def build_compare(slide) -> None:
    add_title(slide, "总结 · 早期与当前", "从“大任务 + 事后处理”转为“分片 + 即时验证”", title_size=26)
    rows = [
        ("任务组织", "较大的固定任务", "短分片，自适应规模"),
        ("失败恢复", "重跑较大输出段", "只补缺失分片"),
        ("质量控制", "批次结束后集中检查", "每片完成立即验收"),
        ("分析入口", "分析直接读取大文件", "一次解析为 compact catalogs"),
        ("保存方式", "原始与派生结果耦合", "原始层 + receipt + 派生层"),
        ("汇总边界", "以批次为主要单位", "几何 / 模式 / 粒子族严格分域"),
    ]
    table_shape = slide.shapes.add_table(
        len(rows) + 1,
        3,
        Inches(0.84),
        Inches(1.52),
        Inches(11.65),
        Inches(4.42),
    )
    table = table_shape.table
    widths = (1.72, 4.33, 5.60)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for col, header in enumerate(("环节", "早期流程", "当前流程")):
        set_cell(table.cell(0, col), header, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER, fill=HEADER_BLUE)
    for row_idx, row in enumerate(rows, start=1):
        fill = WHITE if row_idx % 2 else PANEL
        set_cell(table.cell(row_idx, 0), row[0], size=10.5, color=BLUE, bold=True, align=PP_ALIGN.CENTER, fill=fill)
        set_cell(table.cell(row_idx, 1), row[1], size=10.5, color=MUTED, fill=fill)
        set_cell(table.cell(row_idx, 2), row[2], size=10.5, color=INK, bold=True, fill=fill)
    add_text(
        slide,
        0.82,
        6.25,
        11.69,
        0.37,
        "下一步：用代表轨道状态做 pilot，实测成本后再扩展完整轨道链。",
        size=14,
        color=GREEN,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = "从气球到近地轨道：入射源、计算负担与当前模拟流程"
    prs.core_properties.subject = "Mass_model_511 / S3d-O8 简化技术报告"
    prs.core_properties.author = "TES-511 project"
    prs.core_properties.keywords = "TES-511, balloon, LEO, incident source, simulation workflow"
    blank = prs.slide_layouts[6]

    builders = [
        build_cover,
        lambda s: build_source_slide(
            s,
            section="01 · 环境对比",
            title="入射源：两类环境不是整体缩放",
            image="source_components.png",
            caption="轨道侧拆分天区 γ、反照成分以及原初/次级带电粒子；气球侧保留八粒子族与 20 个角箱。",
        ),
        lambda s: build_source_slide(
            s,
            section="01 · 关键能带",
            title="轨道环境并非全面更高或更低",
            image="gamma_source.png",
            caption="γ：100–300 keV 为 1.17×，300 keV–10 MeV 为 0.30–0.60×；1–10 MeV 质子为 21.5×、正电子为 4.55×。",
        ),
        build_burden,
        build_pipeline,
        build_storage,
        build_compare,
    ]

    for index, (builder, note) in enumerate(zip(builders, NOTES)):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE
        add_chrome(slide, index, len(builders))
        builder(slide)
        add_notes(slide, note[0], note[1])

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build()
    print(f"wrote {output} ({output.stat().st_size} bytes)")
