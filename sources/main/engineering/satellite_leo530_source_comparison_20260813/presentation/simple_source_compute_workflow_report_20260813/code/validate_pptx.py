#!/usr/bin/env python3
"""Fail-closed structural/content validation for the editable briefing PPTX."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "TES511_balloon_to_LEO_source_compute_workflow_brief_20260813.pptx"
REPORT = ROOT / "pptx_validation.json"
EXPECTED_SLIDES = 7

FORBIDDEN = [
    "1000倍",
    "单位错误",
    "能量缩放",
    "Geant4提速",
    "压缩成62.8",
    "卫星本底一定更低",
    "轨道环境更干净",
]

REQUIRED = [
    "入射源级",
    "不含活化",
    "不是探测器计数",
    "6,149,560 validated primaries",
    "488 validated receipts",
    "13,777,092",
    "77.6 GB",
    "62.8 MB",
    "代表轨道状态",
]


def slide_text(slide) -> str:
    values: list[str] = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            values.append(shape.text)
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def main() -> int:
    problems: list[str] = []
    warnings: list[str] = []
    if not PPTX.exists():
        problems.append("PPTX does not exist")
        REPORT.write_text(
            json.dumps({"status": "FAIL", "problems": problems}, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        return 1

    prs = Presentation(PPTX)
    if len(prs.slides) != EXPECTED_SLIDES:
        problems.append(f"expected {EXPECTED_SLIDES} slides, got {len(prs.slides)}")

    expected_size = (12191695, 6858000)
    actual_size = (prs.slide_width, prs.slide_height)
    if actual_size != expected_size:
        problems.append(f"unexpected slide size {actual_size}; expected {expected_size}")

    details: list[dict[str, object]] = []
    all_text: list[str] = []
    note_text: list[str] = []
    for index, slide in enumerate(prs.slides, start=1):
        text = slide_text(slide)
        all_text.append(text)
        pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        tables = sum(getattr(shape, "has_table", False) for shape in slide.shapes)
        text_shapes = sum(getattr(shape, "has_text_frame", False) for shape in slide.shapes)
        out_of_bounds = []
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                out_of_bounds.append(shape.name)
            if shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                out_of_bounds.append(shape.name)
        if out_of_bounds:
            problems.append(f"slide {index} out-of-bounds shapes: {out_of_bounds}")
        if text_shapes < 3:
            problems.append(f"slide {index} has too few editable text shapes ({text_shapes})")
        if index in (2, 3) and pictures != 1:
            problems.append(f"slide {index} expected exactly one embedded figure, got {pictures}")
        if index not in (2, 3) and pictures:
            warnings.append(f"slide {index} unexpectedly contains {pictures} picture(s)")
        if index == 7 and tables != 1:
            problems.append(f"slide 7 expected one editable table, got {tables}")

        note = slide.notes_slide.notes_text_frame.text.strip()
        note_text.append(note)
        if len(note) < 25:
            problems.append(f"slide {index} speaker notes missing or too short")
        details.append(
            {
                "slide": index,
                "pictures": pictures,
                "tables": tables,
                "text_shapes": text_shapes,
                "shape_count": len(slide.shapes),
                "out_of_bounds_shapes": out_of_bounds,
                "notes_chars": len(note),
                "title_preview": re.sub(r"\s+", " ", text)[:100],
            }
        )

    combined = "\n".join(all_text + note_text)
    for token in FORBIDDEN:
        if token in combined:
            problems.append(f"forbidden wording present: {token}")
    for token in REQUIRED:
        if token not in combined:
            problems.append(f"required wording missing: {token}")

    # The batch metric must be called primaries, never detector events.
    if re.search(r"6[,.]?149[,.]?560\s+(?:events|事件)", combined, flags=re.IGNORECASE):
        problems.append("batch0006 primary count is mislabeled as events")

    animations: list[str] = []
    transitions: list[str] = []
    external_relationships: list[str] = []
    with zipfile.ZipFile(PPTX) as archive:
        slide_xml = [
            name
            for name in archive.namelist()
            if name.startswith("ppt/slides/slide") and name.endswith(".xml")
        ]
        animations = [name for name in slide_xml if b"<p:timing" in archive.read(name)]
        transitions = [name for name in slide_xml if b"<p:transition" in archive.read(name)]
        external_relationships = [
            name
            for name in archive.namelist()
            if name.endswith(".rels") and b'TargetMode="External"' in archive.read(name)
        ]
    if animations:
        problems.append(f"animations present: {animations}")
    if transitions:
        problems.append(f"transitions present: {transitions}")
    if external_relationships:
        problems.append(f"external relationships present: {external_relationships}")

    report = {
        "status": "PASS" if not problems else "FAIL",
        "pptx": str(PPTX),
        "bytes": PPTX.stat().st_size,
        "slides": len(prs.slides),
        "slide_size_emu": actual_size,
        "slide_details": details,
        "animations": animations,
        "transitions": transitions,
        "external_relationships": external_relationships,
        "forbidden_tokens_checked": FORBIDDEN,
        "required_tokens_checked": REQUIRED,
        "warnings": warnings,
        "problems": problems,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not problems else 1


if __name__ == "__main__":
    raise SystemExit(main())
