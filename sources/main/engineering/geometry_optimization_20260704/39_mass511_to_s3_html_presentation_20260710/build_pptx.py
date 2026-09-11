#!/usr/bin/env python3
"""Build an editable white-background PowerPoint version of the review deck."""

from __future__ import annotations

from math import ceil
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


WORK = Path(__file__).resolve().parent
ASSETS = WORK / "assets"
OUTPUT = WORK / "Mass_511_to_S3_optimization_review_20260710.pptx"

WIDE_W = 13.333
WIDE_H = 7.5

FONT = "Microsoft YaHei"
INK = RGBColor(23, 32, 51)
MUTED = RGBColor(82, 97, 116)
DIM = RGBColor(115, 132, 151)
BORDER = RGBColor(220, 227, 237)
PANEL = RGBColor(251, 252, 254)
CALLOUT = RGBColor(244, 247, 251)
BLUE = RGBColor(40, 85, 217)
TEAL = RGBColor(20, 125, 115)
ORANGE = RGBColor(237, 104, 70)
PURPLE = RGBColor(120, 55, 238)
AMBER = RGBColor(228, 167, 45)
GREY = RGBColor(108, 123, 145)


def asset(name: str) -> Path:
    path = ASSETS / name
    if not path.exists():
        raise FileNotFoundError(path)
    return path


SLIDES = [
    {
        "image": "s3_2d_processed.png",
        "kicker": "TES 511 keV balloon · project review",
        "title": "Mass_511 到 S3",
        "lead": "几何、本底与中子屏蔽复盘。用同一组数据回答结构改了什么、本底下降多少，以及含硼聚乙烯对中子和 O/Cu 活化有多大作用。",
        "metrics": [("−70.0%", "511 keV 线窗总本底", BLUE), ("−46.3%", "内区相互作用中子率", TEAL)],
        "source": "2026-07-10 · 白底静态版 · 所有图来自项目保留结果或本次审计重算",
    },
    {
        "image": "plain_performance.png",
        "kicker": "结论先行",
        "title": "S3 已把重点从“屏蔽不完整”转成“如何压低残余来源”",
        "bullets": [
            "总本底：0.07412 → 0.02223 次/秒，下降 70.0%。",
            "20 天、3σ 线通量阈值：5.24 → 2.85 ×10⁻⁵ ph cm⁻² s⁻¹，改善 45.6%。",
            "剩余本底：大气 511 keV 线 48.9%，正电子 36.6%。",
            "派生方案：先闭环 S3a；质量预算足够时再评估 S3c。",
        ],
        "callout": "中子结论：慢化很明显，但快中子尾没有消失，不能只看中位能量。",
        "source": "线窗：510.58–511.42 keV；S3 为当前直接探测器响应，完整任务时间轴仍待重建。",
    },
    {
        "image": "mass_2d_processed.png",
        "kicker": "起点 · Mass_511",
        "title": "真实的 45° 低温机基准，但主动屏蔽仍有缺口",
        "lead": "Mass_511 把分级冷板、支撑和侧向开口放入同一质量模型，是后续材料与包覆实验的母几何。",
        "metrics": [
            ("180.5 kg", "基准质量模型", GREY),
            ("0.07412/s", "匹配后线窗总本底", ORANGE),
            ("5.24×10⁻⁵", "20 天、3σ 阈值", PURPLE),
            ("45°", "低温机与屏蔽真实姿态", TEAL),
        ],
        "callout": "它提供了可信基准；问题是主动拒斥层不连续，中子层和底部高 Z 层也未形成明确分工。",
        "source": "来源：Mass_model_511 完成审计，并加入同一强度假设的大气 511 keV 线模拟。",
    },
    {
        "image": "s1_2d_processed.png",
        "kicker": "第一轮屏蔽",
        "title": "把带电粒子、中子和底部伽马分给不同材料处理",
        "bullets": [
            "5 mm 塑料闪烁体：主动标记穿过外层的带电粒子。",
            "10 mm、5 wt% 含硼聚乙烯：氢负责减速，硼负责吸收慢中子。",
            "5 mm 底部钨：只处理底部方向高能光子，不把高 Z 材料铺满全壳。",
        ],
        "metrics": [("−24.2%", "总本底，相对 Mass_511", TEAL), ("−24.9%", "固定活度，整套层栈对比", BLUE)],
        "callout": "证据边界：这是整套层栈对比，不是只开关含硼聚乙烯的纯 A/B 试验。",
        "source": "总本底 0.05617 次/秒；固定活度 141.83 → 106.48 Bq。",
    },
    {
        "image": "s2b_2d_processed.png",
        "kicker": "几何过渡 · S2b",
        "title": "让外层壳真正跟随 45° 低温机，而不是局部补丁",
        "bullets": [
            "含硼聚乙烯增至 20 mm，外层塑闪增至 10 mm。",
            "壳体围绕仪器坐标构建，保留侧向信号窗和顶部服务开口。",
            "为 S3 的完整主动晶体壳提供正确外边界。",
        ],
        "metrics": [("+53.7 kg", "外层几何与材料代价", GREY), ("中间站", "没有完整性能闭环", ORANGE)],
        "callout": "S2b 证明的是几何可实现性；性能结论仍应从 S3 的完整链条读取。",
        "source": "来源：S2b 45° cryostat-following shell 几何审计。",
    },
    {
        "image": "s3_2d_processed.png",
        "kicker": "主线版本 · S3",
        "title": "关键跃迁：用 4 cm 主动晶体连续包住低温机侧壁与端部",
        "lead": "S3 保留 20 mm 含硼聚乙烯与 10 mm 外层塑闪，把局部 CsI 改成完整主动壳，同时保留必要开口。",
        "metrics": [
            ("205.8 kg", "全包覆 CsI 质量", TEAL),
            ("0.02223/s", "匹配后总本底", BLUE),
            ("2.85×10⁻⁵", "20 天、3σ 阈值", PURPLE),
            ("10.52σ", "20 天参考信号显著性", ORANGE),
        ],
        "callout": "S3 是第一个把几何、瞬发本底、延迟源、信号输运和探测器响应串起来的优化版本。",
        "source": "延迟源：每个材料家族 50,000 个空间点、每个点 10⁶ 次抽样；结果文件均核对到 S3 几何。",
    },
    {
        "image": "s3_rz_processed.png",
        "kicker": "S3 层栈",
        "title": "从内向外，每一层只承担一种主要任务",
        "bullets": [
            "40 mm CsI：主动标记进入内区的伽马与带电次级粒子。",
            "0.3 mm Kapton + 8 mm Al：包装与机械支撑。",
            "20 mm 含硼聚乙烯：先减速，再俘获中子。",
            "10 mm 塑闪：最外层主动标记带电粒子。",
        ],
        "callout": "不能封死的地方：负 x 侧信号窗与顶部服务口仍在，也是后续边界计数器应重点记录的位置。",
        "source": "S3a 只换主动晶体；S3b 只换机械壳材料；S3c 同时换两者。",
    },
    {
        "image": "s3a_shell_processed.png",
        "kicker": "派生方案 · S3a",
        "title": "把 CsI 换成 BGO：主动晶体材料是最强的单一杠杆",
        "bullets": [
            "拓扑和厚度不变：4 cm CsI → 4 cm BGO。",
            "BGO 质量约 325.35 kg，明显高于 S3 的 CsI。",
            "三类主导本底快速筛选为 0.00514 次/秒，仅为 S3 的 24.4%。",
            "中子专用延迟源固定活度 31.31 Bq，三种派生方案中最低。",
        ],
        "callout": "建议：先把 S3a 接入共同探测器响应和任务时间轴，再谈晋级。",
        "source": "快速筛选只含正电子、中子和大气 511 keV 主导项，不等同于完整任务链。",
    },
    {
        "image": "s3b_shell_processed.png",
        "kicker": "派生方案 · S3b",
        "title": "只在机械壳加钨：质量代价高，快速本底收益有限",
        "bullets": [
            "主动 CsI 不变；8 mm Al → 2 mm W + 3 mm Al。",
            "新增约 53.88 kg W，同时保留约 11.39 kg Al。",
            "主导本底为 0.01814 次/秒，仍是 S3 的 86.1%。",
            "中子专用延迟源固定活度 99.89 Bq，三种派生方案中最高。",
        ],
        "callout": "判断：单独加钨只做微调，还可能增加活化负担，不宜作为第一优先级。",
        "source": "快速筛选与中子专用延迟源均已通过几何指向和归一化检查。",
    },
    {
        "image": "s3c_shell_processed.png",
        "kicker": "派生方案 · S3c",
        "title": "BGO 与钨同时使用：筛选值最低，但工程代价也最大",
        "bullets": [
            "采用 S3a 的 4 cm BGO，再叠加 S3b 的 2 mm W + 3 mm Al。",
            "主导本底快速筛选为 0.00388 次/秒，是 S3 的 18.4%。",
            "估算 20 天、3σ 阈值为 1.36×10⁻⁵ ph cm⁻² s⁻¹。",
            "中子专用延迟源固定活度 57.82 Bq，高于 S3a。",
        ],
        "callout": "定位：S3c 是质量预算充足时的极限选项，不应绕过完整响应和结构审查直接晋级。",
        "source": "相对 S3a，估算阈值只再改善约 10.5%，需要用工程代价判断是否值得。",
    },
    {
        "image": "plain_performance.png",
        "kicker": "匹配性能",
        "title": "从渐进改善到跃迁：完整主动壳带来最大下降",
        "metrics": [
            ("0.0741/s", "Mass_511 本底", GREY),
            ("0.0562/s", "第一轮屏蔽本底", TEAL),
            ("0.0222/s", "S3 本底", BLUE),
            ("−45.6%", "Mass_511 → S3 通量阈值", PURPLE),
        ],
        "bullets": [
            "三者使用同一大气 511 keV 线强度假设，适合做相对比较。",
            "S3 目前是直接探测器响应；完整任务时间轴会影响绝对阈值。",
        ],
        "source": "匹配总本底包含瞬发项、活化衰变和同一名义强度的大气 511 keV 线。",
    },
    {
        "image": "plain_s3_composition.png",
        "kicker": "S3 剩余本底",
        "title": "活化已不是最大项；外部 511 keV 线与正电子合计 85.5%",
        "bullets": [
            "大气 511 keV 线：0.01088/s，占 48.9%。",
            "正电子：0.00814/s，占 36.6%。",
            "中子：0.00204/s，占 9.2%。",
            "活化衰变：0.00117/s，占 5.3%。",
        ],
        "callout": "优化优先级：先处理外部 511 keV 线和未被主动层标记的正电子；中子仍要盯住快尾和材料活化。",
        "source": "占比按各本底流的物理速率归一，不按最终事件个数直接相除。",
    },
    {
        "image": "s3_w2_final_trajectory_overlay.png",
        "kicker": "最终事件方向",
        "title": "83 个最终幸存事件集中指向低温机侧壁与端部",
        "bullets": [
            "正电子 12 个，中子 3 个，大气 511 keV 56 个，活化衰变 12 个。",
            "外部粒子从真实初始方向投影到有限外包络，再连到 TES 能量质心。",
            "活化衰变从内部产生位置连到 TES。",
        ],
        "callout": "读图限制：这些线是“初始方向弦”，不是粒子在材料中的逐步轨迹，也不是边界穿越 scorer。",
        "source": "图中旧标签 W2 指 510.58–511.42 keV 线窗。",
    },
    {
        "image": "s3_w2_trajectories_by_component.png",
        "kicker": "分来源空间指纹",
        "title": "四类本底的空间来源不同，不能用同一块材料解决",
        "bullets": [
            "正电子：侧向为主，也有底部和不命中有限外包络的事件。",
            "中子：仅 3 个最终事件，只能作诊断，不能据此拟合角分布。",
            "大气 511 keV：形成最广的侧向与底部扇区。",
            "活化：起点集中在内仪器区，应从材料与中子能谱入手。",
        ],
        "callout": "全包覆主动晶体对多方向光子有效；单一底部钨只覆盖其中一部分路径。",
        "source": "跨来源比较使用物理速率；类内方向比例才使用事件计数。",
    },
    {
        "image": "plain_entry_directions.png",
        "kicker": "入口方向",
        "title": "侧面是主要入口；大气 511 keV 仍有四分之一来自底部",
        "metrics": [
            ("75%", "正电子从侧面进入", ORANGE),
            ("100%", "中子侧面；仅 3 个事件", BLUE),
            ("75%", "大气 511 keV 从侧面", PURPLE),
            ("25%", "大气 511 keV 从底部", AMBER),
        ],
        "bullets": [
            "这支持连续侧壁主动壳作为主结构。",
            "底部高 Z 层可作补充，但不是主导解法。",
        ],
        "source": "入口为有限外包络上的直线投影代理；17% 正电子未命中外包络，可能先在外部结构作用。",
    },
    {
        "image": "b10_o_cu_activation_cross_sections.png",
        "kicker": "含硼聚乙烯 · 为什么有效",
        "title": "氢先把中子减速，硼再用很大的热中子截面把它吸收",
        "bullets": [
            "¹⁰B(n,α)：热中子截面约 3845 barn。",
            "⁶³Cu(n,γ)⁶⁴Cu：热中子截面约 4.48 barn；约 578 eV 有 336 barn 共振峰。",
            "快中子活化门槛：Cu 约 10.1–11.0 MeV，O 约 16.7 MeV。",
        ],
        "callout": "关键解释：核截面本身不随 Mass_511 或 S3 改变；两种几何的差别来自到达 O/Cu 的中子数、能谱和位置。",
        "source": "截面：ENDF/B-VIII.1，293.15 K，NNDC；barn = 10⁻²⁴ cm²。",
    },
    {
        "image": "mass_s3_postshield_neutron_count.png",
        "kicker": "含硼层之后还剩多少中子",
        "title": "S3 内区发生作用的中子率下降 46.3%",
        "metrics": [
            ("547.5/s", "Mass_511 内区作用率", GREY),
            ("293.8/s", "S3 内区作用率", BLUE),
            ("10.47%", "Mass_511 生成历史占比", GREY),
            ("5.62%", "S3 生成历史占比", BLUE),
        ],
        "bullets": [
            "每种几何使用 8 次重复、合计 770 万个中子历史。",
            "计数定义为“至少一次在内区发生相互作用的主中子历史”。",
        ],
        "callout": "不能过度归因：项目没有完全相同、只开关 BPE 的 A/B，因此 −46.3% 是 S3 整套屏蔽相对 Mass_511 的净效果。",
        "source": "该指标不是边界通量；不计入无相互作用直接穿过内区的中子。",
    },
    {
        "image": "mass_s3_postshield_neutron_spectrum.png",
        "kicker": "中子慢化与能量分布",
        "title": "大多数中子被明显减速，但十几 MeV 的快尾仍然存在",
        "metrics": [
            ("45.9 → 1.11 keV", "内区首次作用前中位能量", TEAL),
            ("99% → 32%", "首次作用前能量 / 初始能量中位比", BLUE),
            ("3.6% → 69.6%", "首次内区作用前已减速的历史", PURPLE),
            ("16.0 → 21.2/s", ">10.1 MeV 的绝对中子率", ORANGE),
        ],
        "callout": "快尾警告：图的纵轴已改为按能量分箱的绝对通量代理。>10.1 MeV 的绝对率约 16.0 → 21.2/s；S3 更大的主动晶体体积会提高快中子被计入的机会，仍需用真实边界通量复核。",
        "source": "能量取真实初始主中子及其在内区第一次相互作用前的母粒子能量；纵轴不带 cm⁻² 归一化。",
    },
    {
        "image": "mass_s3_neutron_only_o_cu_activation.png",
        "kicker": "O / Cu 活化结果",
        "title": "S3 大幅压低 ⁶⁴Cu 和 ⁶⁶Cu，但快中子产物只小幅下降",
        "bullets": [
            "⁶⁴Cu：17.08 → 3.97 Bq，下降 76.8%。",
            "⁶⁶Cu：4.00 → 0.97 Bq，下降 75.7%。",
            "⁶²Cu：0.373 → 0.320 Bq，只下降 14.2%。",
            "¹⁵O：0.01425 → 0.01018 Bq，下降 28.6%。",
            "⁶¹Cu：0.0651 → 0.0529 Bq，下降 18.7%。",
        ],
        "callout": "解释：热/慢中子相关的铜活化下降很大；依赖十几 MeV 快中子的 ⁶²Cu、¹⁵O 下降较小，与残余快尾一致。",
        "source": "8 次中子专用产额重建；每个文件严格检查一个总时间记录，再除以 8 并传播到 15 天固定活度。",
    },
    {
        "image": "plain_s3abc_screening.png",
        "kicker": "派生方案比较",
        "title": "数据支持先做 S3a：收益接近 S3c，质量与活化风险更容易控制",
        "metrics": [
            ("2.85", "S3 阈值，×10⁻⁵", GREY),
            ("1.52", "S3a 估算阈值", BLUE),
            ("2.66", "S3b 估算阈值", TEAL),
            ("1.36", "S3c 估算阈值", PURPLE),
        ],
        "bullets": [
            "S3a 已取得大部分收益，且中子专用固定活度最低。",
            "S3b 只加钨，收益小且活化负担最高。",
            "S3c 最低，但相对 S3a 的估算阈值只再改善约 10.5%。",
        ],
        "callout": "决策：S3a 作为下一条闭环主线；S3c 作为质量预算充足时的备选。",
        "source": "这是同统计量快速筛选；残余项与信号响应暂按 S3 保持，不是最终晋级结论。",
    },
    {
        "image": "s3a_2d_processed.png",
        "kicker": "结论与下一步",
        "title": "保留 S3 的结构逻辑，先把 S3a 做成可比较、可晋级的完整基线",
        "bullets": [
            "已证明：S3 几何与直接探测器响应闭环；S3a/b/c 的中子专用延迟输运均通过。",
            "待补一：给 S3 与 S3a/b/c 使用同一中子参考样本和同一探测器响应选择。",
            "待补二：在侧窗、顶部、侧壁和底部记录真实第一次边界穿越及能量，复核快中子尾。",
            "待补三：重建完整任务时间轴，并传播大气 511 keV 的低/名义/高三种情景。",
        ],
        "metrics": [("S3a first", "近期主线", TEAL), ("S3c option", "质量预算允许时再评估", PURPLE)],
        "callout": "一句话：S3 解决“有没有完整主动壳”；下一轮解决“用什么材料，并确认快中子尾和工程质量是否值得”。",
        "source": "延迟源继续保留 NUBASE 基态修正、逐文件时间守卫与完整来源记录；任何指向错误几何的运行均作废。",
    },
]

TABLE_SLIDES = [
    {
        "kicker": "TES 本底计数表 · 完整闭环",
        "title": "Mass_511：511 keV 线窗最终 TES 本底构成",
        "subtitle": "最终选择：主动 veto < 50 keV，且通过 side-Compton/FoV。单位：cps；20 天 = 1,728,000 s。",
        "rows": [
            ("PROMPT 正电子", "0.0305299", "52,756", "45 个最终事件；当前 Mass_511 Step05"),
            ("PROMPT 中子", "0.0122099", "21,099", "18 个最终事件；当前 Mass_511 Step05"),
            ("其他 prompt（μ⁺）", "0.00134910", "2,331", "2 个最终事件；保留以使总计精确闭合"),
            ("大气 511 keV", "0.0260449", "45,006", "134 个最终事件；名义 4π sidecar"),
            ("活化 delay", "0.00398460", "6,885", "28 个最终事件；exact-position delayed"),
            ("全本底总计", "0.0741185", "128,077", "与 matched 4π ATM511 总本底一致"),
        ],
        "total_row": 5,
        "callout": "前四个用户关心的主分量之外，只有 0.00134910 cps 的 μ⁺ prompt 项；该行不能省略，否则总计无法闭合。",
        "source": "来源：Mass_model_511 full-stat Step05、Mass_511 matched 4π ATM511 3M sidecar。线窗 510.58–511.42 keV。",
    },
    {
        "kicker": "TES 本底计数表 · 快速筛选",
        "title": "S3c：当前已闭环 prompt / ATM511 TES 本底构成",
        "subtitle": "同一线窗与最终选择。该页只报告已有 TES 计数率的三项；不是 S3c 完整本底闭环。",
        "rows": [
            ("PROMPT 正电子", "0.00135997", "2,350", "2 个最终事件；等统计量 prompt 筛选，低统计"),
            ("PROMPT 中子", "0.00135670", "2,344", "2 个最终事件；低统计，仅作筛选"),
            ("大气 511 keV", "0.00116688", "2,016", "6 个最终事件；名义 4π sidecar，统计误差较大"),
            ("活化 delay", "—", "—", "中子专用固定源活动 57.82 Bq；尚未转成 TES 计数率"),
            ("其他 prompt", "未纳入", "—", "本轮 S3c 筛选仅输运正电子与中子"),
            ("已闭环三项小计（非总计）", "0.00388356", "6,711", "仅 正电子 + 中子 + 大气 511 keV"),
        ],
        "total_row": 5,
        "pending_rows": [3, 4],
        "callout": "不能把 57.82 Bq 活动度当作 TES cps，也不能把三项小计写成完整 S3c 本底；必须先完成 delay 的探测器响应与后续时间轴传播。",
        "source": "来源：S3c dominant backgrounds (e⁺/n/ATM511) 与 clean neutron-only delayed transport。线窗 510.58–511.42 keV。",
    },
]


def set_font(run, size: float, color: RGBColor, bold: bool = False, name: str = FONT) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rpr.set(qn("a:ea"), name)


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float,
    color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0.0,
    space_after: float = 0.0,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = line
        set_font(run, size, color, bold)
    return shape


def add_bullets(slide, x: float, y: float, w: float, h: float, bullets: list[str]) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    for index, item in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(4.5)
        p.level = 0
        run = p.add_run()
        run.text = "• " + item
        set_font(run, 14.1, MUTED)
    return shape


def add_card(slide, x: float, y: float, w: float, h: float, value: str, label: str, color: RGBColor) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = BORDER
    card.line.width = Pt(0.75)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.045), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = color
    stripe.line.fill.background()
    value_size = 20 if len(value) <= 12 else 15.5
    add_text(slide, x + 0.15, y + 0.08, w - 0.22, 0.27, value, size=value_size, color=color, bold=True)
    add_text(slide, x + 0.15, y + h - 0.25, w - 0.22, 0.18, label, size=8.5, color=MUTED)


def add_metrics(slide, x: float, y: float, w: float, metrics: list[tuple[str, str, RGBColor]]) -> float:
    count = len(metrics)
    if count == 1:
        cols, card_h = 1, 0.72
    elif count == 2:
        cols, card_h = 2, 0.72
    else:
        cols, card_h = 2, 0.64
    gap = 0.09
    card_w = (w - gap * (cols - 1)) / cols
    rows = ceil(count / cols)
    for i, (value, label, color) in enumerate(metrics):
        row, col = divmod(i, cols)
        add_card(slide, x + col * (card_w + gap), y + row * (card_h + gap), card_w, card_h, value, label, color)
    return rows * card_h + (rows - 1) * gap


def add_visual(slide, image_path: Path) -> None:
    x, y, w, h = 0.48, 0.70, 5.72, 6.15
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = BORDER
    panel.line.width = Pt(0.75)
    with Image.open(image_path) as image:
        iw, ih = image.size
    pad = 0.16
    available_w, available_h = w - 2 * pad, h - 2 * pad
    scale = min(available_w / iw, available_h / ih)
    picture_w, picture_h = iw * scale, ih * scale
    picture_x = x + (w - picture_w) / 2
    picture_y = y + (h - picture_h) / 2
    slide.shapes.add_picture(str(image_path), Inches(picture_x), Inches(picture_y), width=Inches(picture_w), height=Inches(picture_h))


def title_size(text: str) -> float:
    if len(text) > 36:
        return 23.5
    if len(text) > 27:
        return 26.0
    return 29.0


def add_copy(slide, data: dict) -> None:
    x, w = 6.48, 6.30
    add_text(slide, x, 0.63, w, 0.24, data["kicker"], size=9.3, color=BLUE, bold=True)
    add_text(slide, x, 0.90, w, 0.82, data["title"], size=title_size(data["title"]), color=INK, bold=True)
    cursor = 1.80
    if data.get("lead"):
        add_text(slide, x, cursor, w, 0.62, data["lead"], size=14.7, color=MUTED)
        cursor += 0.66
    if data.get("metrics"):
        metric_h = add_metrics(slide, x, cursor, w, data["metrics"])
        cursor += metric_h + 0.13
    if data.get("bullets"):
        bullet_h = max(0.56, 0.34 * len(data["bullets"]) + 0.08)
        add_bullets(slide, x, cursor, w, bullet_h, data["bullets"])
        cursor += bullet_h + 0.08
    if data.get("callout"):
        callout_h = 0.50 if len(data["callout"]) < 68 else 0.66
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(cursor), Inches(w), Inches(callout_h))
        box.fill.solid()
        box.fill.fore_color.rgb = CALLOUT
        box.line.fill.background()
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(cursor), Inches(0.04), Inches(callout_h))
        accent.fill.solid()
        accent.fill.fore_color.rgb = BLUE
        accent.line.fill.background()
        add_text(slide, x + 0.13, cursor + 0.08, w - 0.22, callout_h - 0.11, data["callout"], size=11.2, color=MUTED)
    add_text(slide, x, 6.87, w, 0.27, data["source"], size=7.6, color=DIM)


def add_chrome(slide, index: int, total: int) -> None:
    base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(WIDE_W), Inches(0.035))
    base.fill.solid()
    base.fill.fore_color.rgb = RGBColor(232, 237, 245)
    base.line.fill.background()
    progress = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(WIDE_W * (index + 1) / total), Inches(0.035))
    progress.fill.solid()
    progress.fill.fore_color.rgb = BLUE
    progress.line.fill.background()
    add_text(slide, 12.35, 7.17, 0.50, 0.16, f"{index + 1} / {total}", size=7.8, color=DIM, align=PP_ALIGN.RIGHT)


def set_cell_text(cell, text: str, *, size: float, color: RGBColor, bold: bool = False, align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    cell.text = ""
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.045)
    cell.margin_bottom = Inches(0.045)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, color, bold)


def add_table_slide(slide, data: dict) -> None:
    add_text(slide, 0.52, 0.62, 12.15, 0.22, data["kicker"], size=9.3, color=BLUE, bold=True)
    add_text(slide, 0.52, 0.90, 12.10, 0.50, data["title"], size=28, color=INK, bold=True)
    add_text(slide, 0.52, 1.43, 12.10, 0.25, data["subtitle"], size=11.0, color=MUTED)

    rows = data["rows"]
    table_shape = slide.shapes.add_table(len(rows) + 1, 4, Inches(0.52), Inches(1.88), Inches(12.25), Inches(3.95))
    table = table_shape.table
    for column, width in zip(table.columns, (2.45, 1.78, 1.75, 6.27)):
        column.width = Inches(width)
    headers = ("来源", "最终 TES 计数率 (cps)", "20 天期望 TES 计数", "状态 / 说明")
    for col, text in enumerate(headers):
        cell = table.cell(0, col)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        set_cell_text(cell, text, size=11.1, color=RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)

    pending = set(data.get("pending_rows", []))
    for row_index, row in enumerate(rows, start=1):
        source, rate, counts, note = row
        is_total = row_index - 1 == data["total_row"]
        is_pending = row_index - 1 in pending
        fill = RGBColor(235, 242, 255) if is_total else (RGBColor(255, 249, 239) if is_pending else RGBColor(255, 255, 255))
        ink = INK if not is_pending else MUTED
        for col, value in enumerate((source, rate, counts, note)):
            cell = table.cell(row_index, col)
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
            set_cell_text(
                cell,
                value,
                size=10.6 if col == 3 else 11.2,
                color=ink,
                bold=is_total,
                align=PP_ALIGN.CENTER if col in (1, 2) else PP_ALIGN.LEFT,
            )

    callout = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.52), Inches(6.08), Inches(12.25), Inches(0.52))
    callout.fill.solid()
    callout.fill.fore_color.rgb = CALLOUT
    callout.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.52), Inches(6.08), Inches(0.05), Inches(0.52))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    add_text(slide, 0.68, 6.20, 12.00, 0.26, data["callout"], size=10.5, color=MUTED)
    add_text(slide, 0.52, 6.83, 12.20, 0.20, data["source"], size=8.0, color=DIM)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(WIDE_W)
    prs.slide_height = Inches(WIDE_H)
    prs.core_properties.title = "Mass_511 到 S3｜几何、本底与中子屏蔽复盘"
    prs.core_properties.subject = "Mass_511 到 S3 及 S3a/S3b/S3c 的几何优化、背景与中子活化审计"
    prs.core_properties.author = "TES 511 Balloon project"
    blank = prs.slide_layouts[6]
    total = len(SLIDES) + len(TABLE_SLIDES)
    for index, data in enumerate(SLIDES):
        slide = prs.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(255, 255, 255)
        add_chrome(slide, index, total)
        add_visual(slide, asset(data["image"]))
        add_copy(slide, data)
    for table_index, data in enumerate(TABLE_SLIDES, start=len(SLIDES)):
        slide = prs.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(255, 255, 255)
        add_chrome(slide, table_index, total)
        add_table_slide(slide, data)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build()
    print(f"wrote {output} ({output.stat().st_size} bytes; {len(SLIDES) + len(TABLE_SLIDES)} slides)")
