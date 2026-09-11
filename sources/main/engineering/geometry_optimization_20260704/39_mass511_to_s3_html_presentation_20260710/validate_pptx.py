#!/usr/bin/env python3
"""Structural validation for the editable PowerPoint deliverable."""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


WORK = Path(__file__).resolve().parent
PPTX = WORK / "Mass_511_to_S3_optimization_review_20260710.pptx"
REPORT = WORK / "pptx_validation.json"
EXPECTED_SLIDES = 23


def main() -> int:
    problems: list[str] = []
    if not PPTX.exists():
        problems.append("PPTX output does not exist")
        report = {"status": "FAIL", "pptx": str(PPTX), "problems": problems}
        REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 1

    presentation = Presentation(PPTX)
    if len(presentation.slides) != EXPECTED_SLIDES:
        problems.append(f"expected {EXPECTED_SLIDES} slides, got {len(presentation.slides)}")
    expected_size = (12191695, 6858000)  # 13.333 × 7.5 in in EMUs, rounded by python-pptx
    actual_size = (presentation.slide_width, presentation.slide_height)
    if actual_size != expected_size:
        problems.append(f"unexpected slide size {actual_size}, expected {expected_size}")

    slides: list[dict[str, int]] = []
    for number, slide in enumerate(presentation.slides, start=1):
        pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        tables = sum(getattr(shape, "has_table", False) for shape in slide.shapes)
        text_shapes = sum(getattr(shape, "has_text_frame", False) for shape in slide.shapes)
        out_of_bounds = 0
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                out_of_bounds += 1
            if shape.left + shape.width > presentation.slide_width or shape.top + shape.height > presentation.slide_height:
                out_of_bounds += 1
        slides.append(
            {
                "slide": number,
                "pictures": pictures,
                "tables": tables,
                "text_shapes": text_shapes,
                "shapes": len(slide.shapes),
                "out_of_bounds_shapes": out_of_bounds,
            }
        )
        is_table_slide = number > 21
        if not is_table_slide and pictures < 1:
            problems.append(f"slide {number} has no embedded picture")
        if is_table_slide and tables < 1:
            problems.append(f"table slide {number} has no editable table")
        if text_shapes < 4:
            problems.append(f"slide {number} has too few editable text shapes ({text_shapes})")
        if out_of_bounds:
            problems.append(f"slide {number} has {out_of_bounds} out-of-bounds shapes")

    with zipfile.ZipFile(PPTX) as archive:
        slide_xml = [name for name in archive.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
        timing = [name for name in slide_xml if b"<p:timing" in archive.read(name)]
        transitions = [name for name in slide_xml if b"<p:transition" in archive.read(name)]
        external_links = [name for name in archive.namelist() if name.endswith(".rels") and b'TargetMode="External"' in archive.read(name)]
    if timing:
        problems.append(f"animations present: {timing}")
    if transitions:
        problems.append(f"transitions present: {transitions}")
    if external_links:
        problems.append(f"external relationships present: {external_links}")

    report = {
        "status": "PASS" if not problems else "FAIL",
        "pptx": str(PPTX),
        "bytes": PPTX.stat().st_size,
        "slides": len(presentation.slides),
        "slide_size_emu": actual_size,
        "slide_details": slides,
        "animations": timing,
        "transitions": transitions,
        "external_relationships": external_links,
        "problems": problems,
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not problems else 1


if __name__ == "__main__":
    raise SystemExit(main())
