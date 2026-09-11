#!/usr/bin/env python3
"""Update the existing five-slide PPT0821 deck with SH3 + L2 results."""

from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[4]
PACKAGE = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PPT0821/ppt.pptx"
FIGURES = PACKAGE / "outputs/figures"
OUTPUT = ROOT / "PPT0821/ppt_SH3_L2_20260821.pptx"
SUMMARY = PACKAGE / "outputs/summary.json"


def replace_visual(slide, old_shape, image_path: Path) -> None:
    tree = slide.shapes._spTree
    old_element = old_shape._element
    index = tree.index(old_element)
    box_left, box_top, box_width, box_height = old_shape.left, old_shape.top, old_shape.width, old_shape.height
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    box_ratio = box_width / box_height
    if image_ratio >= box_ratio:
        width = box_width
        height = int(round(width / image_ratio))
        left = box_left
        top = box_top + (box_height - height) // 2
    else:
        height = box_height
        width = int(round(height * image_ratio))
        top = box_top
        left = box_left + (box_width - width) // 2
    tree.remove(old_element)
    picture = slide.shapes.add_picture(str(image_path), left, top, width, height)
    picture_element = picture._element
    tree.remove(picture_element)
    tree.insert(index, picture_element)


def set_text(shape, text: str, *, size_pt: float, bold_first: bool = False, color: str = "222222") -> None:
    frame = shape.text_frame
    frame.clear()
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size_pt)
        run.font.bold = bool(bold_first and index == 0)
        run.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_after = Pt(1.5)


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def main() -> int:
    required = {
        "slide2": FIGURES / "02_full_spectrum_comparison_sh3_l2.png",
        "slide3": FIGURES / "01_background_energy_bands_sh3_l2.png",
        "slide4": FIGURES / "03_normalized_minimum_detectable_flux_sh3_l2.png",
    }
    if not SOURCE.is_file():
        raise RuntimeError(f"Missing source PPT: {SOURCE}")
    for path in required.values():
        if not path.is_file():
            raise RuntimeError(f"Missing replacement figure: {path}")
    if not SUMMARY.is_file():
        raise RuntimeError(f"Missing projection summary: {SUMMARY}")

    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    l2 = summary["main_L2_result"]
    l2_flux = summary["L2_source_contract"]["integrated_flux_over_fixed_model_support_cm2_s"]

    deck = Presentation(SOURCE)
    if len(deck.slides) != 5:
        raise RuntimeError(f"Unexpected source slide count: {len(deck.slides)}")

    # Slide 1: retain the original title-slide treatment.
    deck.slides[0].shapes[0].text = "20260821"

    # Slides 2-4: the first placeholder contains the existing raster figure.
    replace_visual(deck.slides[1], deck.slides[1].shapes[0], required["slide2"])
    set_text(
        deck.slides[1].shapes[1],
        "大气：38 km，34°N，100°E\n"
        "LEO：530 km，近赤道，非 SAA 静态谱\n"
        "月面：Apollo 17，REDMoon\n"
        "L2：近 1 AU；静态 GCR + 宇宙 γ",
        size_pt=11.5,
    )
    set_text(
        deck.slides[1].shapes[2],
        "SH3_OPTV3_60cm 当前基线：20 天、3σ，Fmin = (2.23 ± 0.21)×10⁻⁵ ph cm⁻² s⁻¹；Aeff = 15.09 cm²。\n"
        "L2 与 LEO 共用 COSI 母谱：去 12.6 GV 地磁传输并扩展至 4π；不含反照、SAA、SEP 和 Z>2 重离子。\n"
        "L2 支撑域积分 [cm⁻² s⁻¹]："
        f"γ {l2_flux['gamma']['integrated_flux_over_model_support_cm2_s']:.2f}，"
        f"p {l2_flux['p']['integrated_flux_over_model_support_cm2_s']:.2f}，"
        f"α {l2_flux['alpha']['integrated_flux_over_model_support_cm2_s']:.3f}，"
        f"e⁻ {l2_flux['eminus']['integrated_flux_over_model_support_cm2_s']:.3f}，"
        f"e⁺ {l2_flux['eplus']['integrated_flux_over_model_support_cm2_s']:.4f}。",
        size_pt=10.8,
    )

    replace_visual(deck.slides[2], deck.slides[2].shapes[0], required["slide3"])
    set_text(
        deck.slides[2].shapes[1],
        "SH3 末级本底反溯：\n"
        "prompt 12 个；\n"
        "delayed 111 个。\n\n"
        "活化按粒子族 × 体积 × 母核态对应。\n"
        "阴影为响应加权 10–90% 初级能区。",
        size_pt=11.5,
        bold_first=True,
    )

    replace_visual(deck.slides[3], deck.slides[3].shapes[0], required["slide4"])
    set_text(
        deck.slides[3].shapes[1],
        "L2 中心值\n"
        f"{float(l2['mission_relative_F3_to_SH3_balloon_including_signal_transmission']):.2f}× 气球\n"
        f"Fmin ≈ ({float(l2['estimated_20d_F3_ph_cm2_s']) * 1e5:.2f} ± "
        f"{float(l2['estimated_20d_F3_conditional_MC_sigma_ph_cm2_s']) * 1e5:.2f})×10⁻⁵\n\n"
        f"delayed 占 {100 * float(l2['delayed_fraction_of_projected_background']):.1f}%\n"
        f"p 占 delayed 的 {100 * float(l2['proton_fraction_of_delayed']):.1f}%\n"
        f"5 个末级事件贡献 {100 * float(l2['top_5_delayed_event_fraction']):.1f}%。\n\n"
        "误差：有限响应样本条件统计。\n"
        "未含模型系统误差；非匹配输运。",
        size_pt=10.8,
    )
    # The old floating reference box overlaps the new chart title and repeats
    # values now printed directly on the chart.
    remove_shape(deck.slides[3].shapes[2])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
