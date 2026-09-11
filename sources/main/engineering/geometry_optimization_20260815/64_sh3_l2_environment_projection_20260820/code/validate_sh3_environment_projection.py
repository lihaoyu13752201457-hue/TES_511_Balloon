#!/usr/bin/env python3
"""Small-file validation for the SH3 four-environment projection and PPT."""

from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from PIL import Image
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[4]
PACKAGE = Path(__file__).resolve().parents[1]
TABLES = PACKAGE / "outputs/tables"
FIGURES = PACKAGE / "outputs/figures"
PPT = ROOT / "PPT0821/ppt_SH3_L2_20260821.pptx"
OUT = PACKAGE / "data/validation.json"


def rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def main() -> int:
    errors: list[str] = []
    warnings: list[str] = []
    summary_path = PACKAGE / "outputs/summary.json"
    scan_path = PACKAGE / "data/activation_primary_energy_scan_summary.json"
    if not summary_path.is_file() or not scan_path.is_file():
        raise RuntimeError("Projection or activation summary is missing")
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    scan = json.loads(scan_path.read_text(encoding="utf-8"))
    if summary.get("status") != "PASS__SH3_FOUR_ENVIRONMENT_SOURCE_SPECTRUM_PROJECTION":
        errors.append("projection status is not PASS")
    if scan.get("status") != "PASS__SH3_SELECTED_ACTIVATION_KEYS_EXACT_PRIMARY_ENERGY_JOIN":
        errors.append("activation primary-energy status is not PASS")
    if scan.get("scan_policy") != "ONE_SEMANTIC_PASS__NO_SIM_HASH":
        errors.append("activation scan policy differs from no-hash contract")
    closure = summary.get("SH3_event_closure", {})
    if closure.get("prompt_W2_final_events") != 12 or closure.get("delayed_W2_final_events") != 111:
        errors.append("SH3 W2-final event closure differs from 12 prompt + 111 delayed")
    if not math.isclose(float(closure.get("total_direct_no_coincidence_rate_cps", -1)), 0.009418680739114394, rel_tol=0, abs_tol=2e-15):
        errors.append("SH3 direct W2 rate closure failed")

    performance = rows(TABLES / "sh3_environment_performance_estimates.csv")
    if len(performance) != 4:
        errors.append("performance table does not contain four environments")
    by_env = {row["environment"]: row for row in performance}
    required_envs = {
        "balloon_38km",
        "leo530_quiet_proxy",
        "lunar_surface_proxy",
        "sun_earth_l2_quiet_1au_proxy",
    }
    if set(by_env) != required_envs:
        errors.append("performance environment set differs")
    elif float(by_env["sun_earth_l2_quiet_1au_proxy"]["estimated_20d_F3_ph_cm2_s"]) <= 0:
        errors.append("L2 F3 is non-positive")
    for environment, row in by_env.items():
        if float(row["total_rate_cps"]) < 0 or float(row["estimated_20d_F3_ph_cm2_s"]) <= 0:
            errors.append(f"non-physical performance value: {environment}")

    bands = rows(TABLES / "sh3_response_weighted_primary_energy_bands.csv")
    if len(bands) != 5:
        errors.append("response-band table does not contain five plotted families")
    for row in bands:
        p10 = float(row["response_weighted_primary_energy_p10_MeV"])
        p50 = float(row["response_weighted_primary_energy_p50_MeV"])
        p90 = float(row["response_weighted_primary_energy_p90_MeV"])
        if not (0 < p10 <= p50 <= p90):
            errors.append(f"non-monotonic response band: {row['family']}")

    expected_figures = [
        FIGURES / "01_background_energy_bands_sh3_l2.png",
        FIGURES / "02_full_spectrum_comparison_sh3_l2.png",
        FIGURES / "03_normalized_minimum_detectable_flux_sh3_l2.png",
    ]
    for path in expected_figures:
        if not path.is_file():
            errors.append(f"missing figure: {path.name}")
            continue
        with Image.open(path) as image:
            if image.width < 1500 or image.height < 900:
                errors.append(f"figure resolution is too small: {path.name} {image.size}")

    if not PPT.is_file():
        errors.append("updated PPT is missing")
    else:
        deck = Presentation(PPT)
        if len(deck.slides) != 5:
            errors.append("updated PPT no longer has five slides")
        text = "\n".join(
            shape.text
            for slide in deck.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        for token in ("SH3_OPTV3_60cm", "L2", "20260821"):
            if token not in text:
                errors.append(f"updated PPT missing token: {token}")
        for stale in ("2.1 e-5 cps", "比较三种不同环境"):
            if stale in text:
                errors.append(f"updated PPT retains stale text: {stale}")

    limitations = summary.get("cautions", [])
    if not any("heavy ions" in item for item in limitations):
        warnings.append("L2 heavy-ion limitation is not explicit")
    if not any("transport" in item for item in limitations):
        warnings.append("matched-transport limitation is not explicit")

    status = "PASS" if not errors else "FAIL"
    payload = {
        "schema_version": 1,
        "status": status,
        "errors": errors,
        "warnings": warnings,
        "checks": {
            "no_large_file_hashing": True,
            "four_environment_rows": len(performance),
            "response_band_rows": len(bands),
            "ppt_slides": len(Presentation(PPT).slides) if PPT.is_file() else 0,
            "figures": len([path for path in expected_figures if path.is_file()]),
        },
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps(payload, indent=2, ensure_ascii=False))
    return 0 if status == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
