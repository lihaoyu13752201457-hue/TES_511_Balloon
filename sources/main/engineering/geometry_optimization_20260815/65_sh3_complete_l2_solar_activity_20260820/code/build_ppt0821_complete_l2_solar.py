#!/usr/bin/env python3
"""Update the concise PPT0821 deck with five-family L2 and solar scenarios."""

from __future__ import annotations

import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[4]
PACKAGE = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "PPT0821/ppt.pptx"
FIGURES = PACKAGE / "outputs/figures"
SUMMARY = PACKAGE / "outputs/summary.json"
OUTPUT = ROOT / "PPT0821/ppt_SH3_L2_five_family_solar_20260821.pptx"


def fitted_box(image_path: Path, left: int, top: int, width: int, height: int) -> tuple[int, int, int, int]:
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio >= box_ratio:
        fitted_width = width
        fitted_height = int(round(fitted_width / image_ratio))
        return left, top + (height - fitted_height) // 2, fitted_width, fitted_height
    fitted_height = height
    fitted_width = int(round(fitted_height * image_ratio))
    return left + (width - fitted_width) // 2, top, fitted_width, fitted_height


def replace_visual(slide, old_shape, image_path: Path) -> None:
    tree = slide.shapes._spTree
    old_element = old_shape._element
    index = tree.index(old_element)
    box = fitted_box(image_path, old_shape.left, old_shape.top, old_shape.width, old_shape.height)
    tree.remove(old_element)
    picture = slide.shapes.add_picture(str(image_path), *box)
    picture_element = picture._element
    tree.remove(picture_element)
    tree.insert(index, picture_element)


def add_visual(slide, image_path: Path, left: int, top: int, width: int, height: int) -> None:
    slide.shapes.add_picture(str(image_path), *fitted_box(image_path, left, top, width, height))


def set_text(
    shape,
    text: str,
    *,
    size_pt: float,
    bold_first: bool = False,
    color: str = "222222",
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Pt(3)
    frame.margin_right = Pt(3)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(1.5)
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size_pt)
        run.font.bold = bool(bold_first and index == 0)
        run.font.color.rgb = RGBColor.from_string(color)


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def move_last_slide_before(deck: Presentation, before_index: int) -> None:
    slide_ids = deck.slides._sldIdLst
    last = slide_ids[-1]
    slide_ids.remove(last)
    slide_ids.insert(before_index, last)


def main() -> int:
    required = {
        "spectra": FIGURES / "02_full_spectrum_comparison_sh3_l2.png",
        "response": FIGURES / "01_background_energy_bands_sh3_l2.png",
        "performance": FIGURES / "03_normalized_minimum_detectable_flux_sh3_l2.png",
        "solar": FIGURES / "04_l2_solar_activity_sh3.png",
    }
    if not SOURCE.is_file() or not SUMMARY.is_file():
        raise RuntimeError("Missing source PPT or L2 summary")
    for path in required.values():
        if not path.is_file():
            raise RuntimeError(f"Missing figure: {path}")

    summary = json.loads(SUMMARY.read_text(encoding="utf-8"))
    full = summary["main_L2_result_solar_min_2009"]
    solar_max = summary["L2_solar_activity_result"]["solar_max_2014"]
    dominant = summary["dominant_L2_activation_key"]

    deck = Presentation(SOURCE)
    if len(deck.slides) != 5:
        raise RuntimeError(f"Unexpected source slide count: {len(deck.slides)}")

    deck.slides[0].shapes[0].text = "20260821"

    replace_visual(deck.slides[1], deck.slides[1].shapes[0], required["spectra"])
    set_text(
        deck.slides[1].shapes[1],
        "L2 quiet 五族（4pi）：\n"
        "p  Athena/Usoskin\n"
        "alpha  Kuznetsov + COSI 接续\n"
        "e-  Athena/Grimani\n"
        "e+  AMS 表外为覆盖代理\n"
        "gamma  COSI + Tuerler",
        size_pt=10.8,
        bold_first=True,
    )
    set_text(
        deck.slides[1].shapes[2],
        "统一口径：单个初级粒子总动能；曲线为声明角域积分 dF/dE。L2 2009 full GCR 最大取 phi=0.3793 GV。\n"
        "外部 neutron/muon 在 L2 物理缺失；HZE 确实存在但 SH3 无响应。软质子、SEP 与方向性银河 gamma 独立，不并入 4pi quiet 基线。",
        size_pt=10.4,
    )

    replace_visual(deck.slides[2], deck.slides[2].shapes[0], required["response"])
    set_text(
        deck.slides[2].shapes[1],
        "SH3 末级本底反溯\n"
        "prompt 12；delayed 111\n\n"
        "活化按粒子族 × 体积 × 母核态对应。\n"
        "prompt 蓝；delayed 橙\n"
        "阴影：各自响应加权 10–90% 初级能区。\n\n"
        "L2 结果主要受 1.37 GeV 附近质子活化控制。",
        size_pt=10.3,
        bold_first=True,
    )

    replace_visual(deck.slides[3], deck.slides[3].shapes[0], required["performance"])
    set_text(
        deck.slides[3].shapes[1],
        "L2 2009 full GCR 最大\n"
        f"{float(full['mission_relative_F3_to_SH3_balloon_including_signal_transmission']):.2f}x 气球\n"
        f"Fmin = ({float(full['estimated_20d_F3_ph_cm2_s']) * 1e5:.2f} +/- "
        f"{float(full['estimated_20d_F3_conditional_MC_sigma_ph_cm2_s']) * 1e5:.2f})e-5\n\n"
        f"{int(dominant['shared_delayed_survivors'])} 个 Cu-62 survivor 共用 "
        f"{int(dominant['activation_RP_records'])} 个 {float(dominant['primary_energy_MeV']) / 1000.0:.3f} GeV p RP；\n"
        f"占 L2 总本底 {100 * float(dominant['fraction_of_total_L2_background']):.1f}%。\n\n"
        "误差仅为条件 MC；不是 L2 最终性能，需匹配输运确认。",
        size_pt=10.1,
        bold_first=True,
    )
    remove_shape(deck.slides[3].shapes[2])

    solar_slide = deck.slides.add_slide(deck.slide_layouts[6])
    add_visual(solar_slide, required["solar"], 257933, 720000, 9470048, 5000000)
    solar_text = solar_slide.shapes.add_textbox(9858133, 577987, 1810264, 5909310)
    set_text(
        solar_text,
        "quiet-GCR 五族代理\n"
        f"2009 full：{float(full['mission_relative_F3_to_SH3_balloon_including_signal_transmission']):.2f}x\n"
        f"Fmin {float(full['estimated_20d_F3_ph_cm2_s']) * 1e5:.2f}e-5\n\n"
        f"2014：{float(solar_max['mission_relative_F3_to_SH3_balloon_including_signal_transmission']):.2f}x\n"
        f"Fmin {float(solar_max['estimated_20d_F3_ph_cm2_s']) * 1e5:.2f}e-5\n\n"
        "跨周期散布约 26%；\n"
        "SEP/HZE/软质子及单点 RP\n"
        "相关误差未计。",
        size_pt=10.0,
        bold_first=True,
    )
    move_last_slide_before(deck, 4)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    deck.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
