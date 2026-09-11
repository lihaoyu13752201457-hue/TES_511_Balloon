#!/usr/bin/env python3
"""Build a readable NIMA-centered PPT with Phase08-12 as extensions."""

from __future__ import annotations

import csv
import json
import shutil
from datetime import datetime
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
R2 = ROOT / "reports2.0"
OUT = R2 / "PPT_FULL_PROGRESS_REPORT"

PRIMARY_PPTX = OUT / "cosmosray_bg_2605_nima_foundation_phase08_12_extension_report.pptx"
LEGACY_PPTX = OUT / "cosmosray_bg_2605_phase07_12_progress_report.pptx"
TOP_MEMORY = R2 / "memoryforPPT.md"
LOCAL_MEMORY = OUT / "memoryforPPT.md"
MD_REPORT = OUT / "phase07_12_progress_report.md"
MANIFEST = OUT / "ppt_manifest.csv"

SW = Inches(13.333333)
SH = Inches(7.5)
FONT = "Microsoft YaHei"

NAVY = RGBColor(20, 42, 68)
BLUE = RGBColor(0, 92, 175)
GREEN = RGBColor(27, 125, 96)
ORANGE = RGBColor(188, 102, 28)
RED = RGBColor(173, 52, 45)
TEXT = RGBColor(34, 41, 47)
MUTED = RGBColor(93, 103, 120)
LIGHT = RGBColor(246, 248, 251)
LINE = RGBColor(215, 222, 232)
WHITE = RGBColor(255, 255, 255)


def r2(rel: str) -> Path:
    return R2 / rel


def read_json(rel: str, default=None):
    p = r2(rel)
    if not p.exists():
        return default
    return json.loads(p.read_text(encoding="utf-8"))


def read_csv(rel: str) -> list[dict[str, str]]:
    p = r2(rel)
    if not p.exists():
        return []
    with p.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def first_row(rel: str) -> dict[str, str]:
    rows = read_csv(rel)
    return rows[0] if rows else {}


def get_lineage(lineage: list[dict], key: str, default=""):
    for row in lineage:
        if row.get("quantity") == key:
            return row.get("value", default)
    return default


def fnum(v, nd=6) -> str:
    if v is None:
        return ""
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    ax = abs(x)
    if ax != 0 and (ax < 1e-3 or ax >= 1e5):
        return f"{x:.{nd}e}"
    return f"{x:.{nd}g}"


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def context() -> dict:
    lineage = read_json("07_NIMA_MANUSCRIPT/final_numerical_lineage.json", [])
    phase10 = read_json("10_POINT_DIFFUSE_DISCRIMINATION/phase10_summary.json", {})
    phase11 = read_json("11_METRIC_RECONCILIATION_AND_OPTICS_GATE/phase11_summary.json", {})
    phase12 = read_json("12_FINAL_COMPACT_SOURCE_ANALYSIS/phase12_summary.json", {})
    metric_rows = read_csv("12_FINAL_COMPACT_SOURCE_ANALYSIS/metric_closure_final.csv")
    metric = {r.get("metric_id", ""): r for r in metric_rows}
    sel11 = read_json("11_METRIC_RECONCILIATION_AND_OPTICS_GATE/selection_best_upgrade_decision.json", {})
    sel12 = read_json("12_FINAL_COMPACT_SOURCE_ANALYSIS/selection_upgrade_decision_final.json", {})
    source9 = read_json("09_SOURCE_CASES_ABC/source_case_summary.json", {})
    fp = read_json(
        "12_FINAL_COMPACT_SOURCE_ANALYSIS/first_principles_channeling_optics/"
        "first_principles_optics_coupling_summary.json",
        {},
    )
    fov_rows = read_csv(
        "12_FINAL_COMPACT_SOURCE_ANALYSIS/no_direct_scaling_optics_alignment/"
        "tables/fov_footprint_summary.csv"
    )
    fov = {r.get("metric_id", ""): r.get("value", "") for r in fov_rows}
    aeff_total = {}
    for row in read_csv(
        "12_FINAL_COMPACT_SOURCE_ANALYSIS/no_direct_scaling_optics_alignment/"
        "tables/aeff_transport_debug.csv"
    ):
        if row.get("ring_id") == "TOTAL":
            aeff_total = row
            break
    return {
        "lineage": lineage,
        "phase10": phase10,
        "phase11": phase11,
        "phase12": phase12,
        "metric": metric,
        "sel11": sel11,
        "sel12": sel12,
        "source9": source9,
        "fp": fp,
        "fov": fov,
        "aeff_total": aeff_total,
        "design_best": first_row(
            "08_DESIGN_OPTIMIZATION_ADDON/WP_D1_selection_pareto_highstat/"
            "top20_catalog_uniform_highstat.csv"
        ),
        "phase10_selection": next(
            (
                r
                for r in read_csv("10_POINT_DIFFUSE_DISCRIMINATION/selection_best_measured_energy_audit.csv")
                if r.get("selection_id") == "bgo30_r18_cent_single_top1_L5_keep"
                and r.get("energy_basis") == "measured"
                and r.get("window") == "broad_480_550"
            ),
            {},
        ),
    }


NIMA_FIGURES: list[dict[str, str]] = [
    {
        "id": "NIMA Fig.1",
        "title": "Science source authority：Be-window source placement",
        "path": "03_NEXT_PHASE_SUPPORT/gate_A_source_placement/first_tes_hit_z_hist.png",
        "takeaway": "Gate-A 修正把 511 keV source 固定为 Be-window HomogeneousBeam；response authority 从这里开始。",
        "note": "避免把旧 source placement 或真实天体源与 detector-response source 混用。",
    },
    {
        "id": "NIMA Fig.1b",
        "title": "511 keV source 与 line broadening 入口",
        "path": "04_FIGURES/nima_update/science_source_and_line_broadening.png",
        "takeaway": "论文主线先定义 science stream，再讨论 intrinsic line model 对窄窗/宽窗灵敏度的影响。",
        "note": "当前主 authority 仍是 mono/reference source；展宽是模型扫描。",
    },
    {
        "id": "NIMA Fig.1c",
        "title": "光学也会接受 FoV/bandpass 内 γ 本底",
        "path": "04_FIGURES/nima_update/optics_focused_gamma_background_concept.png",
        "takeaway": "NIMA 论文已经把 optics-focused gamma background 作为单独 Level-1 aperture addendum 引入。",
        "note": "这不是 production optics；它是防止漏算/双计数的 ledger 补项。",
    },
    {
        "id": "NIMA Fig.1d",
        "title": "Optics-focused γ addendum 的量级",
        "path": "04_FIGURES/nima_update/optics_focused_gamma_background_addendum.png",
        "takeaway": "480-550 keV addendum 约 1.02e-6 cps，对当前阈值影响极小，但必须显式记录。",
        "note": "小项也进入 final numerical lineage。",
    },
    {
        "id": "NIMA Fig.2",
        "title": "Prompt cosmic-ray profile：粒子 scale 随时间变化",
        "path": "04_FIGURES/phase2/phase2_particle_scale_by_day.png",
        "takeaway": "NIMA 的环境输入来自 Phase2 reference profile，而不是单一常数本底。",
        "note": "workspace validation 保留 prompt lightcurve static WARN，论文中必须作为限制。",
    },
    {
        "id": "NIMA Fig.3",
        "title": "Prompt final rate 的 day-scale 曲线",
        "path": "04_FIGURES/phase2/phase2_prompt_rate_day_curve.png",
        "takeaway": "prompt component 被映射到共同时间轴，用于后续 broad-window ledger。",
        "note": "这是 reference profile，不是实测 telemetry。",
    },
    {
        "id": "NIMA Fig.4",
        "title": "511 keV atmospheric transmission",
        "path": "04_FIGURES/phase2/phase2_science_transmission.png",
        "takeaway": "science response 需要同时经过 Aeff、Tatm 和 visibility/geometry 因子。",
        "note": "当前 corrected response 是 24.858994 cps/flux。",
    },
    {
        "id": "NIMA Fig.5a",
        "title": "Delayed source 的 RPIP sampling 证据",
        "path": "03_NEXT_PHASE_SUPPORT/activation_rpip_sampling/plot_check_day15_2602.png",
        "takeaway": "活化源位置来自 buildup isotope-production positions，不是均匀撒点。",
        "note": "这是 delayed-source 方法学的关键基础图。",
    },
    {
        "id": "NIMA Fig.5b",
        "title": "RPIP sampling 与 XZTES 几何上下文",
        "path": "04_FIGURES/nima_update/rpip_sampling_geometry_context.png",
        "takeaway": "抽样点和输运几何在同一质量模型中闭合。",
        "note": "用于支撑 delayed activation spatial model。",
    },
    {
        "id": "NIMA Fig.5c",
        "title": "Activation geometry volume map",
        "path": "04_FIGURES/nima_update/activation_geometry_volume_map.png",
        "takeaway": "几何体、activity 和 transported delayed rate 被放在同一张图中审计。",
        "note": "这张图是 NIMA 版本比基础报告更完整的证据。",
    },
    {
        "id": "NIMA Fig.5",
        "title": "W183/W180 修正后的 activation Top10",
        "path": "04_FIGURES/day15/activation_top10_after_fix.png",
        "takeaway": "ground-state 修正后，day-15 activation contributor 排序重新冻结。",
        "note": "这是 delayed source fix 后的 authority，不回退旧图。",
    },
    {
        "id": "NIMA Fig.6",
        "title": "Delayed 511 keV 区域 top nuclide 能谱",
        "path": "03_NEXT_PHASE_SUPPORT/activation_511_diagnostics/delayed_511_energy_spectrum_by_top_nuclides.png",
        "takeaway": "511 keV 附近 delayed background 的谱形来自核素组合，而不是单条线。",
        "note": "支撑 delayed-dominated 的物理解释。",
    },
    {
        "id": "NIMA Fig.7",
        "title": "Delayed final rate 的核素贡献",
        "path": "03_NEXT_PHASE_SUPPORT/activation_511_diagnostics/delayed_511_top10_bar.png",
        "takeaway": "延迟活化是 broad 480-550 keV final background 的主导来源。",
        "note": "后续优化优先处理 delayed 与材料/几何。",
    },
    {
        "id": "NIMA Fig.8",
        "title": "480-550 keV 共同时间轴 VETO chain 谱",
        "path": "04_FIGURES/day15/timeline_spectrum_480_550_veto_chain.png",
        "takeaway": "NIMA 论文的核心结果图：同一时间轴下 raw/BGO/final 逐级收敛。",
        "note": "该图应优先出现在汇报主干，而不是附录。",
    },
    {
        "id": "NIMA Fig.9",
        "title": "raw、BGO、final 三阶段 rate 对比",
        "path": "04_FIGURES/day15/timeline_veto_rates_bar.png",
        "takeaway": "BGO 与 final cuts 明显压低本底，但 delayed 仍留下主导残余。",
        "note": "这是从 0416 基础报告继承并更新的主结果叙事。",
    },
    {
        "id": "NIMA Fig.10",
        "title": "IMAGE8-like component spectrum with science stream",
        "path": "04_FIGURES/day15/image8_like_component_spectrum_with_science.png",
        "takeaway": "prompt、delayed、science diagnostic stream 在同一谱图中对齐。",
        "note": "science stream 仅 diagnostic，不加入 background-only sensitivity ledger。",
    },
    {
        "id": "NIMA Fig.11",
        "title": "Broad window true/measured spectrum",
        "path": "03_NEXT_PHASE_SUPPORT/gate_B_detector_response/spectrum_480_550_true_vs_measured.png",
        "takeaway": "Gate-B detector response 把 true-energy 与 measured-energy 差异显式化。",
        "note": "后续 Phase10/12 measured-energy audit 依赖这条线。",
    },
    {
        "id": "NIMA Fig.12",
        "title": "Line region true/measured spectrum",
        "path": "03_NEXT_PHASE_SUPPORT/gate_B_detector_response/spectrum_510_515_true_vs_measured.png",
        "takeaway": "窄线窗口对 detector response 更敏感，不能只看 true-energy 统计。",
        "note": "也是 Phase12 measured ERL 的前置动机。",
    },
    {
        "id": "NIMA Fig.13",
        "title": "PARMA scale outlier contribution audit",
        "path": "01B_CONVERGENCE_PATCH_UPDATE/parma_outlier_audit/parma_scale_outlier_heatmap.png",
        "takeaway": "高 scale bin 的 contribution-weighted 影响经过审计；uncapped baseline 被保留为小系统误差。",
        "note": "避免把环境 scale 极值误判成主导不确定性。",
    },
    {
        "id": "NIMA Fig.14",
        "title": "Line-model sensitivity vs intrinsic width",
        "path": "03_NEXT_PHASE_SUPPORT/science_line_models/sensitivity_by_line_model.png",
        "takeaway": "intrinsic line width 会改变 line/broad window sensitivity；论文需要明确模型依赖。",
        "note": "不能把 mono-line threshold 泛化到所有天体源谱型。",
    },
    {
        "id": "NIMA Fig.15",
        "title": "不同 line models 的 true/measured spectra",
        "path": "03_NEXT_PHASE_SUPPORT/science_line_models/spectrum_true_vs_measured_by_model.png",
        "takeaway": "source spectrum 与 measured response 的卷积会改变窗口内可见信号。",
        "note": "这正是后续 A/B/C 源型层的入口。",
    },
    {
        "id": "NIMA Fig.16",
        "title": "IMAGE8-style broad-window component spectrum",
        "path": "04_FIGURES/phase2/phase2_image8_broad_480_550.png",
        "takeaway": "Phase2/NIMA broad-window component decomposition 是背景论文的基础表述。",
        "note": "作为 baseline Ta6 validated reference 保留。",
    },
    {
        "id": "NIMA Fig.17",
        "title": "IMAGE8-style line-window component spectrum",
        "path": "04_FIGURES/phase2/phase2_image8_line_510p3_511p8.png",
        "takeaway": "line-window 比 broad-window 更接近 511 keV 目标，但更依赖 detector response 和核素线结构。",
        "note": "不能只用 line-window 给最终 science claim。",
    },
    {
        "id": "NIMA Fig.18",
        "title": "Top nuclide activity evolution",
        "path": "04_FIGURES/phase2/phase2_top_nuclide_activity_vs_day.png",
        "takeaway": "活化随时间积累/衰减，day-15 是一个被选择并审计的 reference point。",
        "note": "后续 exposure/trajectory 变化会重新改变 delayed ledger。",
    },
    {
        "id": "NIMA Fig.19",
        "title": "Fisher 与 profiled nuisance threshold 对比",
        "path": "04_FIGURES/phase2/phase2_profiled_vs_fisher_threshold.png",
        "takeaway": "论文灵敏度不是简单 sqrt(B)；profiled nuisance proxy 给出更现实的阈值。",
        "note": "也是 Phase11 metric crosswalk 的原因。",
    },
    {
        "id": "NIMA Fig.20",
        "title": "P(>=3σ) vs flux",
        "path": "04_FIGURES/phase2/phase2_P3_vs_flux_profiled.png",
        "takeaway": "F=1e-4、1 Ms 在当前 proxy 下不是稳健 3σ。",
        "note": "这是 claim-control 的核心图。",
    },
    {
        "id": "NIMA Fig.21",
        "title": "Template likelihood vs window counting",
        "path": "03_NEXT_PHASE_SUPPORT/likelihood_511/likelihood_vs_window_counting.png",
        "takeaway": "template likelihood 提供比 window counting 更强的分辨力，但必须声明 proxy/metric 限制。",
        "note": "Phase10 count-only 与 Phase9/12 ERL 不可直接比较。",
    },
    {
        "id": "NIMA Fig.22",
        "title": "Long-timeline injection recovery",
        "path": "03_NEXT_PHASE_SUPPORT/long_timeline_injection/recovered_flux_vs_true_flux.png",
        "takeaway": "注入恢复用于检查统计 proxy 是否有系统偏差。",
        "note": "仍不是最终 real-flight claim。",
    },
    {
        "id": "NIMA Fig.22b",
        "title": "Small statistical dashboard",
        "path": "04_FIGURES/nima_update/nima_small_stat_dashboard.png",
        "takeaway": "把 component rates、thresholds、injection probability 和 delayed activity 放到一页审计。",
        "note": "适合作为 NIMA 基础报告的总览图。",
    },
    {
        "id": "NIMA Fig.23",
        "title": "Science survival by timing model",
        "path": "04_FIGURES/phase2/phase2_science_survival_by_timing_model.png",
        "takeaway": "timing/DAQ cuts 会影响 science survival，不能只优化 background。",
        "note": "Phase08 的 timing overlay 不能直接与 catalog rows 混合。",
    },
    {
        "id": "NIMA Fig.24",
        "title": "Background rate vs coincidence window",
        "path": "03_NEXT_PHASE_SUPPORT/timing_window_scan/background_rate_vs_window.png",
        "takeaway": "coincidence window 缩短会降低 background，但也会改变 signal survival。",
        "note": "下一步需要真实 electronics/DAQ 支撑。",
    },
    {
        "id": "NIMA Fig.25",
        "title": "Science survival vs coincidence window",
        "path": "03_NEXT_PHASE_SUPPORT/timing_window_scan/science_survival_vs_window.png",
        "takeaway": "时间窗优化必须用 signal/background 双指标；不能只看 residual background。",
        "note": "这给 Phase08/12 requirements 留出电子学边界。",
    },
    {
        "id": "NIMA Fig.26",
        "title": "Phase08 high-stat Pareto：Q/Q0 vs background",
        "path": "08_DESIGN_OPTIMIZATION_ADDON/WP_D1_selection_pareto_highstat/highstat_q_vs_background.png",
        "takeaway": "NIMA 论文基础已经引入 design add-on：selection-only best 有潜力，但只是候选。",
        "note": "后续 Phase10/12 负责审计它能否升级。",
    },
    {
        "id": "NIMA Fig.27",
        "title": "Phase08 source acceptance vs background",
        "path": "08_DESIGN_OPTIMIZATION_ADDON/WP_D1_selection_pareto_highstat/highstat_source_acceptance_vs_background.png",
        "takeaway": "降低 background 往往伴随 source acceptance 损失；这就是为什么不能只按 Q 排名。",
        "note": "selection-only best 后续未升级为主结果。",
    },
]


EXTENSION_FIGURES: list[dict[str, str]] = [
    {
        "phase": "Phase09",
        "title": "A compact-source detectability folding",
        "path": "09_SOURCE_CASES_ABC/figures/A_GC_POINT_detectability.png",
        "takeaway": "A 是主 compact-source case；它继承 NIMA 的 response/background authority。",
        "note": "association_status 只能是 assumption，不是 confirmed source identity。",
    },
    {
        "phase": "Phase09",
        "title": "B diffuse aperture foreground",
        "path": "09_SOURCE_CASES_ABC/figures/B_diffuse_expected_line_background.png",
        "takeaway": "B 只作为 FoV aperture foreground/null model，不生成 focal-spot Cosima source。",
        "note": "这修正了“把弥散总 flux 当点源输运”的风险。",
    },
    {
        "phase": "Phase10",
        "title": "A vs B detection probability scaffold",
        "path": "10_POINT_DIFFUSE_DISCRIMINATION/figures/A_vs_B_detection_probability.png",
        "takeaway": "Phase10 给 point/diffuse discrimination 一个 L1 scaffold，但仍是 placeholder optics。",
        "note": "count-only metric 是保守诊断，不替代 NIMA ERL proxy。",
    },
    {
        "phase": "Phase10",
        "title": "Selection best true/measured audit",
        "path": "10_POINT_DIFFUSE_DISCRIMINATION/figures/selection_best_true_vs_measured.png",
        "takeaway": "selection-only best 通过 measured-energy audit，但这不等于升级为 main analysis。",
        "note": "BGO 仍是 event-total proxy。",
    },
    {
        "phase": "Phase11",
        "title": "Metric F3 crosswalk",
        "path": "11_METRIC_RECONCILIATION_AND_OPTICS_GATE/figures/metric_F3_comparison.png",
        "takeaway": "Phase9/10 的主要差异来自 metric definition，不是单一物理失败。",
        "note": "同一页不能混排 ERL/profiled 与 count-only 后直接排名。",
    },
    {
        "phase": "Phase11",
        "title": "Selection upgrade decision tree",
        "path": "11_METRIC_RECONCILIATION_AND_OPTICS_GATE/figures/selection_upgrade_decision_tree.png",
        "takeaway": "audited、reproduced、upgraded 是三种状态；当前只到 audited。",
        "note": "selection-only best 没有替代 baseline。",
    },
    {
        "phase": "Phase11",
        "title": "Production optics schema gate",
        "path": "11_METRIC_RECONCILIATION_AND_OPTICS_GATE/figures/optics_schema_overview.png",
        "takeaway": "没有 production optics authority 时，A/B 判别只能是 schema/scaffold/requirement。",
        "note": "需要 Aeff(E,theta)、PSF、FoV、bandpass、focal map。",
    },
    {
        "phase": "Phase12",
        "title": "Metric closure：Phase9/10/12",
        "path": "12_FINAL_COMPACT_SOURCE_ANALYSIS/final_figures/metric_closure_phase9_phase10_phase12.png",
        "takeaway": "Phase12 baseline measured ERL F3 接近 1e-4，但未通过最终闭环升级条件。",
        "note": "这是最终 Decision C 的关键证据。",
    },
    {
        "phase": "Phase12",
        "title": "Final selection decision matrix",
        "path": "12_FINAL_COMPACT_SOURCE_ANALYSIS/final_figures/selection_decision_matrix.png",
        "takeaway": "主结果保持 baseline；selection-only best 仍是 secondary design comparison。",
        "note": "防止优化表静默替换主分析。",
    },
    {
        "phase": "Phase12",
        "title": "Optics requirement view",
        "path": "12_FINAL_COMPACT_SOURCE_ANALYSIS/final_figures/optics_bandpass_or_requirement.png",
        "takeaway": "最终包收口为 optics/electronics/response requirements，而不是 final detectability claim。",
        "note": "Decision C / PARAMETRIC_OPTICS_REQUIREMENT。",
    },
    {
        "phase": "Optics",
        "title": "First-principles L2 focal spot",
        "path": "12_FINAL_COMPACT_SOURCE_ANALYSIS/first_principles_channeling_optics/figures/first_principles_L2_focal_spot.png",
        "takeaway": "first-principles optics 已接入 Phase12，但只作为 requirement input。",
        "note": "不使用 CAM511 calibration；不替代 production optics。",
    },
    {
        "phase": "Optics",
        "title": "W/Si Parratt reflectivity",
        "path": "12_FINAL_COMPACT_SOURCE_ANALYSIS/first_principles_channeling_optics/figures/reflectivity_wsi_parratt.png",
        "takeaway": "反射权重来自 W/Si Parratt table，不是经验 lookup。",
        "note": "L2 Aeff 明显低于 50.89 cm2 normalization。",
    },
    {
        "phase": "Optics",
        "title": "On-axis core vs full-FoV footprint",
        "path": "12_FINAL_COMPACT_SOURCE_ANALYSIS/no_direct_scaling_optics_alignment/figures/onaxis_core_vs_full_fov_footprint.png",
        "takeaway": "CAM511 few-cm scale 应先与 full-FoV footprint 对齐，不应直接贴 on-axis core。",
        "note": "no-direct-scaling alignment 的主结论。",
    },
    {
        "phase": "Optics",
        "title": "Aeff transport debug",
        "path": "12_FINAL_COMPACT_SOURCE_ANALYSIS/no_direct_scaling_optics_alignment/figures/aeff_by_ring_and_bounce.png",
        "takeaway": "Aeff/open geometric 由 reflectivity product 决定；表中 used_for_tuning=false。",
        "note": "不要用乘系数与 CAM511 做硬拟合。",
    },
]


class Deck:
    def __init__(self, ctx: dict):
        self.ctx = ctx
        self.prs = Presentation()
        self.prs.slide_width = SW
        self.prs.slide_height = SH
        self.assets: list[tuple[str, str]] = []

    def blank(self):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
        bg.fill.solid()
        bg.fill.fore_color.rgb = WHITE
        bg.line.color.rgb = WHITE
        return slide

    def footer(self, slide, section: str):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.18), SW, Inches(0.32))
        bar.fill.solid()
        bar.fill.fore_color.rgb = LIGHT
        bar.line.color.rgb = LINE
        t = slide.shapes.add_textbox(Inches(0.45), Inches(7.23), Inches(12.3), Inches(0.18))
        p = t.text_frame.paragraphs[0]
        p.text = f"NIMA foundation + Phase08-12 extensions | {section}"
        p.font.name = FONT
        p.font.size = Pt(8)
        p.font.color.rgb = MUTED
        p.alignment = PP_ALIGN.RIGHT

    def text(self, slide, x, y, w, h, txt, size=14, color=TEXT, bold=False, align=None):
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = txt
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.bold = bold
        if align:
            p.alignment = align
        return box

    def title(self, slide, title: str, eyebrow: str = "", section: str = ""):
        if eyebrow:
            self.text(slide, Inches(0.55), Inches(0.32), Inches(12), Inches(0.25), eyebrow, 9, BLUE, True)
        self.text(slide, Inches(0.55), Inches(0.58), Inches(12), Inches(0.56), title, 23, NAVY, True)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.56), Inches(1.21), Inches(1.2), Inches(0.05))
        accent.fill.solid()
        accent.fill.fore_color.rgb = BLUE
        accent.line.color.rgb = BLUE
        self.footer(slide, section or eyebrow)

    def bullet_panel(self, slide, x, y, w, h, bullets: list[str], heading: str = "读图结论", color=BLUE):
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        panel.fill.solid()
        panel.fill.fore_color.rgb = LIGHT
        panel.line.color.rgb = LINE
        self.text(slide, x + Inches(0.18), y + Inches(0.15), w - Inches(0.36), Inches(0.28), heading, 11, color, True)
        yy = y + Inches(0.55)
        for b in bullets[:4]:
            self.text(slide, x + Inches(0.22), yy, w - Inches(0.42), Inches(0.6), f"• {b}", 12.2, TEXT)
            yy += Inches(0.68)

    def image(self, slide, path: Path, x, y, w, h, border=True):
        if not path.exists():
            self.text(slide, x, y + h / 2 - Inches(0.2), w, Inches(0.4), f"缺图: {rel(path)}", 13, RED, True)
            return
        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(w / iw, h / ih)
        nw = int(iw * scale)
        nh = int(ih * scale)
        px = x + (w - nw) / 2
        py = y + (h - nh) / 2
        if border:
            rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
            rect.fill.solid()
            rect.fill.fore_color.rgb = WHITE
            rect.line.color.rgb = LINE
        slide.shapes.add_picture(str(path), px, py, width=nw, height=nh)
        self.assets.append((rel(path), "figure used in NIMA-centered PPT"))

    def cover(self):
        slide = self.blank()
        bg = slide.shapes[0]
        bg.fill.fore_color.rgb = NAVY
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), SH)
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = BLUE
        stripe.line.color.rgb = BLUE
        self.text(slide, Inches(0.85), Inches(1.05), Inches(11.9), Inches(0.9), "球载 511 keV TES 谱仪本底模拟与紧致源拓展", 31, WHITE, True)
        self.text(slide, Inches(0.88), Inches(2.12), Inches(11.4), Inches(0.5), "以 NIMA 论文为基础，Phase08-12 作为拓展、修正与收口", 20, RGBColor(218, 231, 247))
        self.text(
            slide,
            Inches(0.92),
            Inches(3.35),
            Inches(10.7),
            Inches(1.5),
            "核心叙事：NIMA 先建立可信的本底/响应/灵敏度谱系；后续 Phase08-12 只在此基础上做设计优化、源型分层、点/弥散判别、指标重对齐和最终 claim-control。",
            16,
            RGBColor(236, 242, 250),
        )
        self.text(slide, Inches(0.92), Inches(6.55), Inches(11.6), Inches(0.25), f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", 9, RGBColor(185, 201, 222))

    def section(self, title: str, subtitle: str, section: str, color=BLUE):
        slide = self.blank()
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.24))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = color
        stripe.line.color.rgb = color
        self.text(slide, Inches(0.85), Inches(2.05), Inches(11.6), Inches(0.38), section, 13, color, True)
        self.text(slide, Inches(0.85), Inches(2.55), Inches(11.6), Inches(0.8), title, 30, NAVY, True)
        self.text(slide, Inches(0.88), Inches(3.55), Inches(10.6), Inches(0.8), subtitle, 16, MUTED)

    def agenda(self):
        lin = self.ctx["lineage"]
        slide = self.blank()
        self.title(slide, "新版 PPT 的结构", "先 NIMA，后拓展", "Deck logic")
        items = [
            ("1", "NIMA 论文基础", "source authority、环境/活化、detector response、VETO chain、统计灵敏度。"),
            ("2", "Phase08-12 拓展", "设计优化、A/B/C 源型、point/diffuse scaffold、metric crosswalk、Decision C。"),
            ("3", "光学收口", "first-principles channeling optics 与 no-direct-scaling alignment 只作为 requirement input。"),
            ("4", "最终口径", "可写结论、禁止结论、下一步 production requirements。"),
        ]
        y = Inches(1.62)
        for num, head, body in items:
            self.callout_row(slide, y, num, head, body)
            y += Inches(1.22)
        self.small_metrics(
            slide,
            [
                ("B 480-550", f"{fnum(get_lineage(lin, 'background-only 480-550'))} cps", GREEN),
                ("Response", f"{fnum(get_lineage(lin, 'corrected science response'))} cps/flux", BLUE),
                ("Final decision", str(self.ctx["phase12"].get("final_decision_code", "C")), ORANGE),
            ],
            y=Inches(6.45),
        )

    def callout_row(self, slide, y, num, head, body):
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.78), y, Inches(0.46), Inches(0.46))
        circ.fill.solid()
        circ.fill.fore_color.rgb = BLUE
        circ.line.color.rgb = BLUE
        self.text(slide, Inches(0.78), y + Inches(0.08), Inches(0.46), Inches(0.2), num, 12, WHITE, True, PP_ALIGN.CENTER)
        self.text(slide, Inches(1.45), y - Inches(0.02), Inches(10.7), Inches(0.34), head, 18, NAVY, True)
        self.text(slide, Inches(1.45), y + Inches(0.44), Inches(10.8), Inches(0.4), body, 13.5, TEXT)

    def small_metrics(self, slide, items: list[tuple[str, str, RGBColor]], y):
        x = Inches(0.75)
        for label, value, color in items:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.8), Inches(0.56))
            box.fill.solid()
            box.fill.fore_color.rgb = LIGHT
            box.line.color.rgb = LINE
            self.text(slide, x + Inches(0.16), y + Inches(0.1), Inches(1.15), Inches(0.2), label, 9.5, MUTED, True)
            self.text(slide, x + Inches(1.28), y + Inches(0.08), Inches(2.35), Inches(0.26), value, 13.5, color, True)
            x += Inches(4.05)

    def figure_slide(self, item: dict, section: str, color=BLUE):
        slide = self.blank()
        self.title(slide, item["title"], item.get("id", item.get("phase", "")), section)
        self.image(slide, r2(item["path"]), Inches(0.58), Inches(1.45), Inches(8.65), Inches(5.45))
        self.bullet_panel(
            slide,
            Inches(9.45),
            Inches(1.58),
            Inches(3.32),
            Inches(4.62),
            [item["takeaway"], item["note"]],
            heading="读图结论",
            color=color,
        )
        self.text(slide, Inches(0.7), Inches(6.92), Inches(12), Inches(0.18), rel(r2(item["path"])), 7.3, MUTED)

    def summary_slide(self):
        lin = self.ctx["lineage"]
        slide = self.blank()
        self.title(slide, "NIMA 基础论文给出的主结论", "NIMA foundation", "NIMA synthesis")
        boxes = [
            ("本底 ledger", f"prompt={fnum(get_lineage(lin, 'prompt final 480-550'))} cps\n"
                         f"delayed={fnum(get_lineage(lin, 'delayed final 480-550'))} cps\n"
                         f"B_total={fnum(get_lineage(lin, 'background-only 480-550'))} cps", GREEN),
            ("响应 authority", f"corrected response\n{fnum(get_lineage(lin, 'corrected science response'))} cps/(ph cm^-2 s^-1)", BLUE),
            ("灵敏度边界", f"ERL F3={fnum(get_lineage(lin, 'broad ERL profiled 3sigma / 1 Ms'))}\n"
                         f"P3@1e-4={get_lineage(lin, 'P>=3sigma at 1e-4, 1 Ms')}", ORANGE),
        ]
        x = Inches(0.72)
        for title, body, color in boxes:
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.68), Inches(3.85), Inches(2.25))
            card.fill.solid()
            card.fill.fore_color.rgb = LIGHT
            card.line.color.rgb = LINE
            self.text(slide, x + Inches(0.22), Inches(1.92), Inches(3.3), Inches(0.3), title, 16, color, True)
            self.text(slide, x + Inches(0.22), Inches(2.42), Inches(3.35), Inches(1.0), body, 15, TEXT)
            x += Inches(4.12)
        self.bullet_panel(
            slide,
            Inches(1.0),
            Inches(4.55),
            Inches(11.35),
            Inches(1.7),
            [
                "论文基础不是“最后一页结论”，而是一套可追溯 authority：source placement、detector response、background ledger、statistical proxy。",
                "后续 Phase08-12 的所有拓展必须在这套 authority 下解释；不能反向改写 NIMA 基础数字。",
            ],
            heading="后续拓展的边界",
            color=BLUE,
        )

    def extension_intro(self):
        slide = self.blank()
        self.title(slide, "Phase08-12 是 NIMA 之后的拓展层", "Extensions", "Phase08-12 map")
        rows = [
            ("08", "设计优化", "找到可能降低本底/提高 Q 的 selection，但不升级主结论。"),
            ("09", "A/B/C 源型", "把主点源、弥散前景/null、V404 benchmark 分层。"),
            ("10", "点/弥散 L1", "保守 count-only scaffold + measured-energy audit。"),
            ("11", "指标重对齐", "解释 Phase9/10 数字差异，建立 production optics gate。"),
            ("12", "终局闭环", "Decision C：requirements-only，不是 final detectability。"),
        ]
        y = Inches(1.45)
        for pid, head, body in rows:
            self.callout_row(slide, y, pid, head, body)
            y += Inches(1.02)

    def decision_slide(self):
        m = self.ctx["metric"]
        p12 = self.ctx["phase12"]
        base = m.get("phase12_baseline_measured_ERL", {})
        sel = m.get("phase12_selection_best_measured_ERL", {})
        slide = self.blank()
        self.title(slide, "Phase12 最终收口", "Decision C", "Final claim")
        self.small_metrics(
            slide,
            [
                ("Final decision", str(p12.get("final_decision_code", "C")), ORANGE),
                ("Primary metric", "baseline measured ERL", BLUE),
                ("Claim level", "PARAMETRIC REQUIREMENT", RED),
            ],
            y=Inches(1.55),
        )
        self.bullet_panel(
            slide,
            Inches(0.85),
            Inches(2.45),
            Inches(5.55),
            Inches(2.65),
            [
                f"baseline measured ERL F3 = {fnum(base.get('F3_1Ms'))}；P3@1e-4 = {fnum(base.get('P_ge_3sigma_at_1e-4_1Ms'))}。",
                f"selection measured ERL F3 = {fnum(sel.get('F3_1Ms'))}，没有优于 baseline 25%。",
                "metric closure、selection upgrade、production optics 三重条件没有同时满足。",
            ],
            heading="结果",
            color=BLUE,
        )
        self.bullet_panel(
            slide,
            Inches(6.75),
            Inches(2.45),
            Inches(5.55),
            Inches(2.65),
            [
                "可以写 requirements-only result。",
                "不能写最终银河中心紧致源探测。",
                "不能写 production optics detectability 或 confirmed source identity。",
            ],
            heading="claim boundary",
            color=RED,
        )

    def optics_numbers_slide(self):
        fp = self.ctx["fp"]
        fov = self.ctx["fov"]
        aeff = self.ctx["aeff_total"]
        slide = self.blank()
        self.title(slide, "光学推进现在能支持什么", "Optics addendum", "Optics conclusion")
        self.small_metrics(
            slide,
            [
                ("L2 Aeff", f"{fnum(fp.get('firstprinciples_aeff_cm2'))} cm2", BLUE),
                ("Scale vs 50.89", fnum(fp.get("response_scale_vs_current_cam511_normalization")), ORANGE),
                ("On-axis D95", f"{fnum(fov.get('intrinsic_onaxis_core_D95_mm'))} mm", GREEN),
            ],
            y=Inches(1.5),
        )
        self.bullet_panel(
            slide,
            Inches(0.85),
            Inches(2.4),
            Inches(5.55),
            Inches(2.75),
            [
                f"FoV radius at 4.47 arcmin / 12 m = {fnum(fov.get('fov_radius_focal_plane_mm'))} mm。",
                f"expected full-FoV envelope D95 = {fnum(fov.get('expected_geometric_full_fov_envelope_D95_mm'))} mm。",
                f"Aeff/open geometric = {fnum(aeff.get('Aeff_over_open_geometric'))}；used_for_tuning=false。",
            ],
            heading="no-direct-scaling 结果",
            color=GREEN,
        )
        self.bullet_panel(
            slide,
            Inches(6.75),
            Inches(2.4),
            Inches(5.55),
            Inches(2.75),
            [
                "first-principles optics 是 requirement input，不是 production response。",
                "CAM511 few-cm spot 应先对齐 metric definition。",
                "不得用直接乘系数把 Aeff 或 spot 贴到 CAM511。",
            ],
            heading="写作边界",
            color=RED,
        )

    def final_slide(self):
        slide = self.blank()
        self.title(slide, "最终可讲的一句话", "Take-home", "Close")
        self.text(
            slide,
            Inches(0.92),
            Inches(1.62),
            Inches(11.8),
            Inches(1.05),
            "NIMA 论文建立了可信的 511 keV TES balloon-background 与 response/sensitivity 基础；Phase08-12 的作用是把设计优化、源型判别和光学需求收进同一 claim-control 框架。",
            21,
            NAVY,
            True,
        )
        self.bullet_panel(
            slide,
            Inches(1.0),
            Inches(3.25),
            Inches(5.5),
            Inches(2.3),
            [
                "当前最强主结果是 Phase12 baseline measured ERL proxy。",
                "结果接近 1e-4，但仍是 requirements-only。",
                "selection 与 optics 都没有升级为 final authority。",
            ],
            heading="可说",
            color=GREEN,
        )
        self.bullet_panel(
            slide,
            Inches(6.85),
            Inches(3.25),
            Inches(5.5),
            Inches(2.3),
            [
                "不声称最终发现。",
                "不声称已确认 source identity。",
                "不声称 production optics detectability。",
            ],
            heading="不可说",
            color=RED,
        )

    def save(self):
        self.prs.save(PRIMARY_PPTX)
        shutil.copy2(PRIMARY_PPTX, LEGACY_PPTX)
        self.assets.append((rel(PRIMARY_PPTX), "primary NIMA-centered PowerPoint report"))
        self.assets.append((rel(LEGACY_PPTX), "legacy filename copy of NIMA-centered PowerPoint report"))


def build_deck(ctx: dict) -> Deck:
    deck = Deck(ctx)
    deck.cover()
    deck.agenda()
    deck.section(
        "NIMA 是基础，不是附录",
        "下面先按 NIMA 论文图序建立 source、background、response、statistics 的主线；Phase08-12 只在这个基础上展开。",
        "NIMA foundation",
        BLUE,
    )
    for item in NIMA_FIGURES:
        deck.figure_slide(item, "NIMA foundation", BLUE)
    deck.summary_slide()
    deck.section(
        "Phase08-12：作为 NIMA 的拓展与收口",
        "这些阶段不替代 NIMA 基础，而是解决 NIMA 后自然出现的问题：selection 能否升级、源型如何分层、指标如何比较、光学缺口如何封口。",
        "Extensions",
        GREEN,
    )
    deck.extension_intro()
    for item in EXTENSION_FIGURES:
        color = GREEN if item.get("phase") in {"Phase09", "Phase10"} else BLUE
        if item.get("phase") == "Phase12":
            color = ORANGE
        if item.get("phase") == "Optics":
            color = RED
        deck.figure_slide(item, item.get("phase", "Extensions"), color)
    deck.decision_slide()
    deck.optics_numbers_slide()
    deck.final_slide()
    return deck


def memory_text(ctx: dict) -> str:
    lin = ctx["lineage"]
    p12 = ctx["phase12"]
    m = ctx["metric"]
    base = m.get("phase12_baseline_measured_ERL", {})
    fp = ctx["fp"]
    fov = ctx["fov"]
    return "\n".join(
        [
            "# memoryforPPT",
            "",
            "本文件记录新版 PPT 的结构记忆：NIMA 是主干，Phase08-12 是拓展、修正和最终收口。",
            "",
            "## 必须保持的叙事结构",
            "",
            "1. 先讲 NIMA 论文基础：source authority、环境/活化本底、VETO chain、detector response、统计灵敏度。",
            "2. 再讲 Phase08-12：设计优化、A/B/C 源型、point/diffuse scaffold、metric crosswalk、Decision C。",
            "3. 最后讲光学：first-principles/no-direct-scaling 只作为 requirements input，不是 production authority。",
            "",
            "## 核心数字",
            "",
            f"- background-only 480-550: `{fnum(get_lineage(lin, 'background-only 480-550'))} cps`.",
            f"- prompt final 480-550: `{fnum(get_lineage(lin, 'prompt final 480-550'))} cps`.",
            f"- delayed final 480-550: `{fnum(get_lineage(lin, 'delayed final 480-550'))} cps`.",
            f"- focused gamma addendum: `{fnum(get_lineage(lin, 'focused gamma broad addendum'))} cps`.",
            f"- corrected science response: `{fnum(get_lineage(lin, 'corrected science response'))} cps/(ph cm^-2 s^-1)`.",
            f"- Phase07 broad ERL F3: `{fnum(get_lineage(lin, 'broad ERL profiled 3sigma / 1 Ms'))}`.",
            f"- Phase12 baseline measured ERL F3: `{fnum(base.get('F3_1Ms'))}`.",
            f"- Phase12 baseline measured ERL P3@1e-4: `{fnum(base.get('P_ge_3sigma_at_1e-4_1Ms'))}`.",
            f"- Phase12 decision: `{p12.get('final_decision_code')}` / `{p12.get('final_claim_level')}`.",
            "",
            "## 光学边界",
            "",
            f"- first-principles L2 Aeff: `{fnum(fp.get('firstprinciples_aeff_cm2'))} cm2`.",
            f"- response scale vs 50.89: `{fnum(fp.get('response_scale_vs_current_cam511_normalization'))}`.",
            f"- intrinsic on-axis D95: `{fnum(fov.get('intrinsic_onaxis_core_D95_mm'))} mm`.",
            f"- expected full-FoV envelope D95: `{fnum(fov.get('expected_geometric_full_fov_envelope_D95_mm'))} mm`.",
            "- 不得把 CAM511 few-cm spot 直接当作 on-axis r95；必须先对齐 metric definition。",
            "",
            "## 输出文件",
            "",
            f"- primary PPT: `{rel(PRIMARY_PPTX)}`",
            f"- compatibility copy: `{rel(LEGACY_PPTX)}`",
            f"- markdown source summary: `{rel(MD_REPORT)}`",
            "",
            "## 更新规则",
            "",
            "- 新增 PPT 页优先复用 NIMA 图或 reports2.0 中已有图件。",
            "- 每页只讲一个判断，少字、大图、清楚标注 claim boundary。",
            "- 如果同一错误重复出现两次，停止试错；找 3-5 种修复方案后实施最高效方案。",
            "",
        ]
    )


def markdown_report(ctx: dict) -> str:
    lin = ctx["lineage"]
    base = ctx["metric"].get("phase12_baseline_measured_ERL", {})
    return "\n".join(
        [
            "# NIMA-centered Phase07-12 PPT source summary",
            "",
            "新版 PPT 以 NIMA 论文为基础，Phase08-12 作为拓展引入。相比旧版，主要改变是：",
            "",
            "- 大幅增加 NIMA 论文图件引用，按论文图序讲清 source、background、response、statistics。",
            "- Phase08-12 不再平铺，而是作为 NIMA 后续问题的回答。",
            "- 每页只保留一个图和 1-2 条读图结论，提高可读性。",
            "",
            "## Key NIMA anchors",
            "",
            f"- Background-only 480-550 keV: `{fnum(get_lineage(lin, 'background-only 480-550'))} cps`.",
            f"- Corrected response: `{fnum(get_lineage(lin, 'corrected science response'))} cps/(ph cm^-2 s^-1)`.",
            f"- Broad ERL F3: `{fnum(get_lineage(lin, 'broad ERL profiled 3sigma / 1 Ms'))}`.",
            f"- P>=3sigma at F=1e-4, 1 Ms: `{get_lineage(lin, 'P>=3sigma at 1e-4, 1 Ms')}`.",
            "",
            "## Phase12 closure",
            "",
            f"- Baseline measured ERL F3: `{fnum(base.get('F3_1Ms'))}`.",
            f"- Baseline measured ERL P3@1e-4: `{fnum(base.get('P_ge_3sigma_at_1e-4_1Ms'))}`.",
            "- Final decision remains Decision C / requirements-only.",
            "",
            "## Figure policy",
            "",
            f"- NIMA figures included: `{len(NIMA_FIGURES)}`.",
            f"- Extension figures included: `{len(EXTENSION_FIGURES)}`.",
            "",
        ]
    )


def copy_scripts():
    dest_dir = R2 / "05_SCRIPTS_AND_CONFIG" / "tools"
    dest_dir.mkdir(parents=True, exist_ok=True)
    for name in ["build_nima_centered_phase07_12_ppt.py", "build_phase07_12_progress_ppt.py"]:
        src = ROOT / "tools" / name
        if src.exists():
            shutil.copy2(src, dest_dir / name)


def write_manifest(deck: Deck):
    rows = []
    for path, role in sorted(set(deck.assets)):
        p = ROOT / path if not Path(path).is_absolute() else Path(path)
        rows.append(
            {
                "relative_path": path,
                "exists": p.exists(),
                "bytes": p.stat().st_size if p.exists() else "",
                "role": role,
            }
        )
    rows.extend(
        [
            {
                "relative_path": rel(TOP_MEMORY),
                "exists": TOP_MEMORY.exists(),
                "bytes": TOP_MEMORY.stat().st_size if TOP_MEMORY.exists() else "",
                "role": "top-level PPT memory context",
            },
            {
                "relative_path": rel(MD_REPORT),
                "exists": MD_REPORT.exists(),
                "bytes": MD_REPORT.stat().st_size if MD_REPORT.exists() else "",
                "role": "NIMA-centered PPT markdown source summary",
            },
        ]
    )
    with MANIFEST.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["relative_path", "exists", "bytes", "role"])
        writer.writeheader()
        writer.writerows(rows)


def update_reports_manifest():
    manifest = R2 / "MANIFEST.tsv"
    lines = []
    for path in sorted(R2.rglob("*")):
        if path.is_file():
            lines.append(f"{path.relative_to(R2)}\t{path.stat().st_size}")
    manifest.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    ctx = context()
    deck = build_deck(ctx)
    TOP_MEMORY.write_text(memory_text(ctx), encoding="utf-8")
    LOCAL_MEMORY.write_text(memory_text(ctx), encoding="utf-8")
    MD_REPORT.write_text(markdown_report(ctx), encoding="utf-8")
    deck.save()
    copy_scripts()
    deck.assets.append((rel(R2 / "05_SCRIPTS_AND_CONFIG/tools/build_nima_centered_phase07_12_ppt.py"), "primary NIMA-centered PPT generator"))
    deck.assets.append((rel(R2 / "05_SCRIPTS_AND_CONFIG/tools/build_phase07_12_progress_ppt.py"), "compatibility wrapper for PPT generation"))
    write_manifest(deck)
    update_reports_manifest()
    print(f"Wrote {PRIMARY_PPTX}")
    print(f"Wrote {LEGACY_PPTX}")
    print(f"Wrote {TOP_MEMORY}")
    print(f"Wrote {MD_REPORT}")
    print(f"Wrote {MANIFEST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
